from datetime import date

from fakes import (
    FakeAutocomplete,
    FakeEmbed,
    FakeFetcher,
    FakeKeywords,
    FakeLLM,
    FakeSearch,
    all_gaps_relevant,
    fake_passage_labels,
)
from seo_engine.brief import write_brief
from seo_engine.deps import Deps
from seo_engine.models import Phrase, Run, TopicCount
from seo_engine.pipeline import run_pipeline
from seo_engine.providers.fetcher import FetchedPage
from seo_engine.providers.tranco import DictRanks

PAGE = """Emitii is a client workspace for agencies.
Share files with clients and collect client approvals in one place.
Clients see project progress without email threads."""

COMPETITOR_TEXT = "\n".join(
    ["File sharing with clients keeps every file in one place for the agency team."] * 5
    + ["Client approvals let clients sign off on work without email back and forth."] * 5
    + ["Time tracking shows hours spent per client project for billing."] * 5
)
URLS = [f"https://p{i}.com/" for i in range(1, 7)]
GOOD_TITLE = "Client Workspace for Agencies: Files & Approvals | Emitii"
DRAFT = {
    "h1": "A client workspace for agencies",
    "intro": "Emitii is a client workspace where agencies share files and collect approvals.",
    "sections": [
        {
            "heading": "Share files with clients",
            "body": "Upload once.\n\n- Latest version\n- Old versions",
        },
        {"heading": "Client approvals", "body": "Clients approve work in one place."},
        {"heading": "Time tracking", "body": "Pricing starts at [ADD: monthly price]."},
    ],
    "faq": [
        {"question": "How do agencies share files with clients?", "answer": "In the workspace."},
        {"question": "Can clients approve work?", "answer": "Yes, on the work itself."},
        {"question": "Is there a free trial?", "answer": "[ADD: trial length]"},
    ],
    "cta": "Start a free workspace for your next client project.",
}
GOOD_DESC = (
    "A client workspace where agencies share files, collect approvals and show progress. "
    "Invite clients in minutes, no email threads."
)


def _llm(draft: dict | list[dict]) -> FakeLLM:
    drafts = draft if isinstance(draft, list) else [draft]

    def next_draft(system: str, user: str) -> dict:
        return drafts.pop(0) if len(drafts) > 1 else drafts[0]

    return FakeLLM(
        {
            "SeedPhrases": lambda s, u: {"phrases": ["client workspace", "agency crm"]},
            "FitVerdicts": lambda s, u: {
                "verdicts": [
                    {"phrase": line[2:], "fits": True}
                    for line in u.split("PHRASES:\n")[1].splitlines()
                ]
            },
            "PageReading": lambda s, u: {
                "page_type": "product",
                "topics": ["file sharing", "client approvals", "time tracking"],
                "questions": [],
            },
            "PassageLabels": fake_passage_labels,
            "GapFits": all_gaps_relevant,
            "NoiseTopics": lambda s, u: {"noise": []},
            "BriefDraft": next_draft,
            "DraftOut": lambda s, u: DRAFT,
        }
    )


def _deps(llm: FakeLLM) -> Deps:
    pages = {
        u: FetchedPage(
            url=u,
            status="ok",
            title="Client workspace",
            text=COMPETITOR_TEXT,
            headings=["Features"],
            word_count=len(COMPETITOR_TEXT.split()),
            method="httpx",
        )
        for u in URLS
    }
    return Deps(
        search=FakeSearch(
            {"client workspace": [(u, "product") for u in URLS]},
            paa={"client workspace": ["How do agencies share files with clients?"]},
        ),
        keywords=FakeKeywords({"client workspace": (40, None), "agency crm": (0, None)}),
        fetcher=FakeFetcher(pages),
        llm=llm,
        embed=FakeEmbed(),
        ranks=DictRanks({}),
        autocomplete=FakeAutocomplete(searched=set()),
    )


def test_pipeline_produces_brief_from_run_state(settings) -> None:
    draft = {
        "titles": [GOOD_TITLE, "Client Workspace Software for Agencies | Emitii"],
        "description": GOOD_DESC,
        "headings": ["Client workspace for agencies", "Time tracking for client projects"],
    }
    run = Run(page_text=PAGE, settings=settings)
    steps: list[tuple[str, str]] = []
    details = run_pipeline(
        run, _deps(_llm(draft)), lambda n, st, d: steps.append((n, st)), today=date(2026, 9, 24)
    )

    assert [n for n, st in steps if st == "done"] == [
        "keywords",
        "serp",
        "competitors",
        "coverage",
        "brief",
        "draft",
    ]
    brief = run.brief
    assert brief is not None
    assert brief.phrases[0].text == "client workspace"
    assert brief.phrases[0].difficulty_source == "computed"
    assert brief.titles[0] == GOOD_TITLE

    # Every count in the brief comes from run state.
    assert all(t in run.coverage for t in brief.must_cover)
    assert all(g in run.gaps for g in brief.gaps)
    assert {t.topic for t in brief.must_cover} >= {
        "file sharing",
        "client approvals",
        "time tracking",
    }
    assert brief.must_cover[0].ours_passages == 0  # never-mentioned topics first

    assert brief.checklist["phrase_in_title"] and brief.checklist["phrase_in_h1"]
    assert brief.checklist["title_width_ok"] and brief.checklist["description_length_ok"]
    assert brief.checklist["intent_matches"]
    assert details.ours_page_type == "product" and len(run.competitors) == 6
    assert details.snippets[0].title == GOOD_TITLE

    draft = brief.draft
    assert draft is not None and draft.h1.startswith("A client workspace")
    assert draft.placeholders == ["[ADD: monthly price]", "[ADD: trial length]"]
    checks = {c.label: c for c in draft.checks}
    assert checks["Main search phrase in the H1"].ok
    assert checks["Main phrase in the first 100 words"].ok
    assert checks["No keyword stuffing"].ok is False  # tiny draft: phrase density is high
    assert not checks["Right length to compete"].ok  # well under 600 words
    assert not checks["Facts to fill in"].ok


def test_writer_rewrites_once_when_every_title_fails(settings) -> None:
    too_long = "The Complete Client Workspace for Marketing Agencies, Studios and Consultancies"
    llm = _llm(
        [
            {"titles": [too_long], "description": GOOD_DESC, "headings": ["Client workspace"]},
            {"titles": [GOOD_TITLE], "description": GOOD_DESC, "headings": ["Client workspace"]},
        ]
    )
    phrase = Phrase(text="client workspace", volume=40, difficulty=10, intent="commercial")
    must = [TopicCount(topic="file sharing", covered_by=6, total=6, ours_passages=0, bucket="must")]
    result = write_brief(llm, PAGE, [phrase], must, [], 0, "", None, [], settings)
    assert result.rewrites == 1
    assert result.brief.titles == [GOOD_TITLE]
    assert llm.calls.count("BriefDraft") == 2


def test_draft_retries_when_h1_misses_the_phrase(settings) -> None:
    from seo_engine.brief import write_draft
    from seo_engine.models import Brief

    bad = {**DRAFT, "h1": "One workspace for agencies", "intro": "Agencies share files here."}
    drafts = [bad, DRAFT]
    llm = FakeLLM({"DraftOut": lambda s, u: drafts.pop(0)})
    phrase = Phrase(text="client workspace", volume=40, difficulty=10, intent="commercial")
    brief = Brief(
        phrases=[phrase],
        titles=[GOOD_TITLE],
        description=GOOD_DESC,
        must_cover=[],
        gaps=[],
        headings=["Client workspace"],
        intent_flag=None,
        score=0,
        checklist={},
    )
    draft = write_draft(llm, PAGE, brief, [], [900, 1200], settings)
    assert llm.calls == ["DraftOut", "DraftOut"]
    assert "client workspace" in draft.h1.lower()


def test_report_docx_reads_like_the_action_plan(settings) -> None:
    import io
    from datetime import datetime

    from docx import Document

    from seo_engine.report import build_report

    run = Run(page_text=PAGE, settings=settings, source_url="https://www.emitii.com/")
    run_pipeline(
        run,
        _deps(
            _llm(
                {
                    "titles": [GOOD_TITLE],
                    "description": GOOD_DESC,
                    "headings": ["Client workspace for agencies"],
                }
            )
        ),
        today=date(2026, 9, 24),
    )
    doc = Document(io.BytesIO(build_report(run, datetime(2026, 9, 24))))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "SEO brief for emitii.com"
    for heading in [
        "1. Target these searches",
        "3. Add these missing topics",
        "6. Suggested content",
        "A client workspace for agencies",
        "Frequently asked questions",
    ]:
        assert heading in texts
    highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
    assert "[ADD: monthly price]" in highlighted
    assert doc.tables[0].rows[1].cells[0].text == "client workspace (main)"
