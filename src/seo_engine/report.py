"""Readable brief report as a Word document (.docx), laid out like the web action plan.

Plain language throughout; the same wording as the web UI (web/src/format.ts).
"""

import io
import re
from datetime import datetime
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from seo_engine.models import Brief, Phrase, Run, TopicCount

BRAND = RGBColor(0x0E, 0x7C, 0x74)
MUTED = RGBColor(0x62, 0x73, 0x7B)
PLACEHOLDER = re.compile(r"(\[ADD:[^\]]*\])")
COUNTRIES = {
    "US": "United States",
    "GB": "United Kingdom",
    "CA": "Canada",
    "AU": "Australia",
    "IN": "India",
    "NP": "Nepal",
    "DE": "Germany",
    "FR": "France",
    "NZ": "New Zealand",
    "IE": "Ireland",
    "SG": "Singapore",
}


def competition(difficulty: int) -> str:
    if difficulty <= 30:
        return "Low"
    return "Medium" if difficulty <= 55 else "High"


def demand(p: Phrase) -> str:
    if p.volume_source == "autocomplete":
        return "People search this on Google"
    if p.volume_source == "bing":
        return f"About {p.volume:,} searches a month (Bing)"
    return f"About {p.volume:,} searches a month"


def intent(value: str) -> str:
    return {
        "commercial": "Comparing options",
        "transactional": "Ready to sign up or buy",
        "informational": "Want to learn",
        "navigational": "Looking for a specific site",
    }.get(value, "Mixed")


def verdict(score: int) -> str:
    return "Strong" if score >= 70 else "Getting there" if score >= 40 else "Needs work"


def site_name(run: Run) -> str | None:
    if not run.source_url:
        return None
    return urlparse(run.source_url).netloc.removeprefix("www.") or None


def key_topics(coverage: list[TopicCount]) -> tuple[int, int]:
    key = [c for c in coverage if c.bucket in ("must", "worth")]
    return sum(1 for c in key if c.ours_passages > 0), len(key)


def _para(
    doc: Document, text: str, *, muted: bool = False, bold: bool = False, size: int | None = None
):
    """Paragraph with [ADD: …] placeholders highlighted in yellow."""
    p = doc.add_paragraph()
    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        run = p.add_run(part)
        run.bold = bold
        if muted:
            run.font.color.rgb = MUTED
        if size:
            run.font.size = Pt(size)
        if PLACEHOLDER.fullmatch(part):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True
    return p


def _body(doc: Document, text: str) -> None:
    """Draft body: blank-line paragraphs, "- " lines as bullets."""
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
        if lines and all(ln.startswith(("- ", "* ", "• ")) for ln in lines):
            for ln in lines:
                _para(doc, ln[2:]).style = doc.styles["List Bullet"]
        elif lines:
            _para(doc, " ".join(lines))


def _table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = text
    doc.add_paragraph()


def _heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = BRAND


def build_report(run: Run, created_at: datetime) -> bytes:
    brief: Brief | None = run.brief
    if brief is None:
        raise ValueError("this run has no brief yet")
    site = site_name(run)
    main = brief.phrases[0]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # Title and summary
    title = doc.add_heading(f"SEO brief for {site or main.text}", level=0)
    for r in title.runs:
        r.font.color.rgb = BRAND
    _para(
        doc,
        f"{created_at:%d %B %Y} · {COUNTRIES.get(run.settings.country, run.settings.country)}"
        f" · {run.settings.site_strength} website"
        + (f" · {run.source_url}" if run.source_url else ""),
        muted=True,
    )

    covered, total = key_topics(run.coverage)
    missing = [t for t in brief.must_cover if t.ours_passages == 0]
    _heading(doc, "Summary", 1)
    _para(doc, f"Score: {brief.score} out of 100 ({verdict(brief.score)})", bold=True, size=13)
    _para(doc, f"Your page covers {covered} of {total} key topics that top-ranking pages share.")
    if brief.intent_flag:
        _para(doc, f"Heads up: {brief.intent_flag}", bold=True)
    _para(doc, "What to do first:", bold=True)
    todos = [f"Aim the page at “{main.text}”, starting with the new title."]
    if missing:
        todos.append(f"Add {len(missing)} topic(s) that most top pages cover.")
    if brief.gaps:
        todos.append(f"Answer {len(brief.gaps)} question(s) searchers ask that competitors skip.")
    if brief.draft and brief.draft.placeholders:
        n = len(brief.draft.placeholders)
        todos.append(f"Fill in the {n} highlighted facts in the suggested content.")
    for t in todos:
        _para(doc, t).style = doc.styles["List Bullet"]

    # 1. Searches
    _heading(doc, "1. Target these searches", 1)
    _para(
        doc, "Searches people really make, where your page has a fair chance to rank.", muted=True
    )
    _table(
        doc,
        ["Search", "Competition", "Demand", "Why people search"],
        [
            [
                p.text + (" (main)" if i == 0 else ""),
                competition(p.difficulty),
                demand(p),
                intent(p.intent),
            ]
            for i, p in enumerate(brief.phrases)
        ],
    )

    # 2. Title and description
    _heading(doc, "2. Update your title and description", 1)
    _para(doc, "Title", bold=True)
    _para(doc, brief.titles[0])
    _para(doc, "Description", bold=True)
    _para(doc, brief.description)
    if len(brief.titles) > 1:
        _para(doc, "Other title ideas", bold=True)
        for t in brief.titles[1:]:
            _para(doc, t).style = doc.styles["List Bullet"]

    # 3. Topics
    _heading(doc, "3. Add these missing topics", 1)
    _para(
        doc,
        "Most pages that rank for your searches cover these. Your page doesn’t mention them yet.",
        muted=True,
    )
    if missing:
        _table(
            doc,
            ["Topic", "Top pages that cover it"],
            [[t.topic.capitalize(), f"{t.covered_by} of {t.total}"] for t in missing],
        )
    else:
        _para(doc, "Your page already covers every must-have topic.")

    # 4. Questions
    _heading(doc, "4. Answer questions others don’t", 1)
    _para(doc, "People search for these, but almost no competing page answers them.", muted=True)
    for g in brief.gaps or []:
        _para(doc, g.topic[:1].upper() + g.topic[1:]).style = doc.styles["List Bullet"]
    if not brief.gaps:
        _para(doc, "No clear unanswered questions this time.")

    # 5. Outline
    _heading(doc, "5. Suggested page outline", 1)
    for i, h in enumerate(brief.headings):
        _para(doc, f"{'Main heading' if i == 0 else 'Section'}: {h}").style = doc.styles[
            "List Bullet"
        ]

    # 6. Suggested content
    if brief.draft:
        d = brief.draft
        _heading(doc, "6. Suggested content", 1)
        _para(
            doc,
            "A draft built only from facts on your page. Highlighted [ADD: …] items need your "
            "input. Read every sentence before publishing.",
            muted=True,
        )
        for c in d.checks:
            _para(
                doc,
                f"{'✓' if c.ok else '✕'} {c.label}" + (f" ({c.detail})" if c.detail else ""),
                muted=c.ok,
            )
        doc.add_paragraph()
        _heading(doc, d.h1, 2)
        _para(doc, d.intro)
        for s in d.sections:
            _heading(doc, s.heading, 3)
            _body(doc, s.body)
        if d.faq:
            _heading(doc, "Frequently asked questions", 3)
            for f in d.faq:
                _para(doc, f.question, bold=True)
                _para(doc, f.answer)
        if d.cta:
            _para(doc, d.cta, bold=True)

    # Appendix
    _heading(doc, "Appendix: pages we compared", 1)
    for c in run.competitors:
        _para(
            doc, f"{c.source.replace('google#', '#')}  {c.url}  ({c.page_type})"
        ).style = doc.styles["List Bullet"]
    _para(doc, f"How the score works: {brief.score_arithmetic}", muted=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
