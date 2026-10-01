"""Every runnable exercise from docs/SEO-Advisor-Guide.md, in one file. Free and offline.

    .venv/bin/python docs/guide_exercises.py                  # list the exercises
    .venv/bin/python docs/guide_exercises.py 4                # run chapter 4's exercises
    .venv/bin/python docs/guide_exercises.py ex09_keyword_research   # run one exercise
    .venv/bin/python docs/guide_exercises.py all              # run every offline exercise
    .venv/bin/python docs/guide_exercises.py check            # check the guide itself
    .venv/bin/python docs/guide_exercises.py sync-guide       # copy exercise code into the guide

How it is laid out:
1. The offline world: a small made-up internet (competitor pages, Google results, Bing
   numbers, a pretend LLM) that the engine runs against instead of the real one.
2. narrate(): prints a finished run step by step.
3. One function per exercise, named exNN_topic, where NN is the chapter.
4. Chapter 17's test file (the test_* functions).
5. The command line and the guide checker.

Exercise output files go to cache/guide_scratch/ (cache/ is git-ignored). Only
ex04_live_run touches the real internet, and it asks before spending anything.
"""

# ruff: noqa: E402, E501, I001
import inspect
import subprocess
# ========================================================================================
# 1. The offline world (read this: every competitor page says what the filters will do)
# ========================================================================================
import re
import sys
from datetime import date
from pathlib import Path

# Make `tests/fakes.py` importable (seo_engine itself is installed by `pip install -e .`).
REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "tests"))

from fakes import (  # noqa: E402
    FakeAutocomplete,
    FakeEmbed,
    FakeFetcher,
    FakeKeywords,
    FakeLLM,
    FakeSearch,
    all_gaps_relevant,
    fake_passage_labels,
)

from seo_engine.config import Settings  # noqa: E402
from seo_engine.deps import Deps  # noqa: E402
from seo_engine.models import Run  # noqa: E402
from seo_engine.providers.fetcher import FetchedPage  # noqa: E402
from seo_engine.providers.tranco import DictRanks  # noqa: E402

TODAY = date(2026, 9, 25)
SCRATCH = REPO / "cache" / "guide_scratch"  # exercise output; cache/ is git-ignored

# ---------------------------------------------------------------------------------------
# 1. Our page (same text as web/src/example.ts)
# ---------------------------------------------------------------------------------------
OUR_PAGE = """Emitii: one workspace for every client project

Agencies lose hours chasing clients across email threads, shared drives and chat apps. Emitii gives each client one shared workspace where the project actually lives.

Share files without the attachment chaos
Upload designs, documents and videos once. Clients always see the latest version, and older versions stay one click away.

Approvals clients actually give
Ask for sign-off on a file or a milestone. Clients approve or leave comments right on the work, and you get a clear record of who approved what and when.

Tasks both sides can see
Keep a shared task list for the project. Assign tasks to your team or to the client, set due dates and see what is blocking progress.

Progress without status meetings
A project timeline shows what is done, what is next and what is waiting on the client, so nobody has to ask for an update.

Built for agencies and studios
Invite clients in minutes. They get a clean, branded space; your team keeps internal notes private.

Start a free workspace for your next client project."""

# ---------------------------------------------------------------------------------------
# 2. Competitor pages. Each paragraph talks about one topic, so you can predict the counts.
# ---------------------------------------------------------------------------------------
PARA = {
    "file sharing": "File sharing with clients keeps every design and document in one place. "
    "Clients download the latest file without digging through email attachments.",
    "client approvals": "Client approvals happen on the work itself. Clients approve a file or "
    "leave comments, and the agency sees who approved each version.",
    "task management": "Task management for client work means a shared task list with owners "
    "and due dates, so both sides know what happens next.",
    "time tracking": "Time tracking records the hours spent on each client project, which "
    "makes billing accurate and shows where the budget goes.",
    "white label branding": "White label branding lets the agency put its own logo and colours "
    "on the portal, so clients see the agency brand, not the vendor.",
    "pricing plans": "Pricing plans usually charge per seat or per client. Most tools offer a "
    "free plan for small teams and paid plans for larger agencies.",
    "data security": "Data security matters when clients upload contracts. Look for encryption, "
    "two-factor sign-in and permissions for every client.",
    "invoices": "Invoices can be sent from the portal. Clients pay invoices online and the "
    "agency sees which invoices are still open.",
    "clientflow integrations": "Clientflow integrations connect Clientflow with Slack, Google "
    "Drive and Zapier, so updates reach the tools your team already uses.",
}


def page_text(*topics: str, repeat: int = 2) -> str:
    """A competitor page: each topic paragraph appears `repeat` times, one paragraph per line."""
    lines = [PARA[t] for t in topics for _ in range(repeat)]
    return "\n".join(lines)


COMPETITORS: dict[str, tuple[str, str]] = {
    # url: (title, text). The comment says what the competitor filters will do with it.
    "https://clientflow.com/": (  # kept
        "Clientflow: Client Portal for Agencies",
        page_text("file sharing", "client approvals", "task management",
                  "clientflow integrations", "pricing plans"),
    ),
    "https://clientflow.com/features": (  # dropped: 3rd page from clientflow.com
        "Clientflow Features",
        page_text("file sharing", "client approvals", "task management", "time tracking"),
    ),
    "https://clientflow.com/pricing": (  # kept (2nd page from clientflow.com)
        "Clientflow Pricing",
        page_text("pricing plans", "file sharing", "invoices", "data security"),
    ),
    "https://portalpro.io/": (  # kept
        "PortalPro | Client Workspace Software",
        page_text("file sharing", "client approvals", "white label branding", "data security"),
    ),
    "https://agencyhub.com/client-portal": (  # kept
        "Agency Client Portal | AgencyHub",
        page_text("file sharing", "task management", "time tracking", "invoices"),
    ),
    "https://sharedspace.app/": (  # kept
        "SharedSpace: Workspace for Client Projects",
        page_text("file sharing", "client approvals", "task management", "pricing plans"),
    ),
    "https://studiodesk.co/": (  # kept
        "StudioDesk | Project Portal for Studios",
        page_text("client approvals", "task management", "white label branding",
                  "file sharing"),
    ),
    "https://filebox.io/clients": (  # kept
        "Filebox for Clients",
        page_text("file sharing", "data security", "client approvals", "pricing plans"),
    ),
    "https://reviewhub.com/best-client-portals": (  # dropped: a listicle, ours is a product
        "12 Best Client Portals for Agencies",
        page_text("file sharing", "pricing plans", "data security", repeat=3),
    ),
    "https://megasuite.com/": (  # dropped: far longer than the others (length outlier)
        "MegaSuite: Everything for Agencies",
        page_text(*PARA, repeat=10),
    ),
    "https://tinyportal.dev/": (  # dropped: too little text to read
        "TinyPortal",
        "A tiny client portal. File sharing with clients.",
    ),
}
FORUM = "https://www.reddit.com/r/agency/comments/1/client_portal"  # dropped: forum
AUTHORITY = "https://www.g2.com/categories/client-portal"  # dropped: authority site


def _fetched(url: str, title: str, text: str) -> FetchedPage:
    words = len(text.split())
    ok = words >= 150  # Thresholds.min_clean_words
    return FetchedPage(
        url=url,
        status="ok" if ok else "too_short",
        title=title,
        text=text,
        headings=[title],
        word_count=words,
        method="httpx",
        reason="" if ok else f"only {words} words of main text",
    )


# ---------------------------------------------------------------------------------------
# 3. Search data: what "Google", "Bing" and "autocomplete" say in this little world
# ---------------------------------------------------------------------------------------
def _serp(*urls: str) -> list[tuple[str, str]]:
    """Google results as (url, page type). The type is what the URL rules would guess."""
    kinds = {FORUM: "forum", AUTHORITY: "category", "https://reviewhub.com/best-client-portals":
             "listicle"}
    return [(u, kinds.get(u, "product")) for u in urls]


SERPS: dict[str, list[tuple[str, str]]] = {
    "client portal for agencies": _serp(
        AUTHORITY,
        "https://clientflow.com/",
        "https://agencyhub.com/client-portal",
        "https://reviewhub.com/best-client-portals",
        "https://megasuite.com/",
        "https://tinyportal.dev/",
    ),
    "client project workspace": _serp(
        "https://portalpro.io/",
        "https://sharedspace.app/",
        "https://studiodesk.co/",
        "https://clientflow.com/features",
        "https://agencyhub.com/client-portal",
        FORUM,
    ),
    "client workspace app": _serp(  # shares 3 results with "client project workspace"
        "https://portalpro.io/",
        "https://sharedspace.app/",
        "https://studiodesk.co/",
    ),
    "file sharing with clients": _serp(
        "https://clientflow.com/",
        "https://clientflow.com/pricing",
        "https://portalpro.io/",
        "https://filebox.io/clients",
        "https://reviewhub.com/best-client-portals",
    ),
    "client portal login": _serp(  # easy to rank for, but the wrong meaning (LLM says no)
        "https://tinyportal.dev/",
        "https://filebox.io/clients",
    ),
    "client approval software": _serp(  # big brands: too hard for a new site
        "https://www.adobe.com/approvals",
        "https://www.salesforce.com/approvals",
        "https://www.microsoft.com/approvals",
        "https://studiodesk.co/",
    ),
    "project management software": _serp(
        "https://www.atlassian.com/",
        "https://www.monday.com/",
        "https://www.asana.com/",
    ),
    "what is a client portal": _serp(
        "https://www.hubspot.com/client-portal",
        "https://www.zendesk.com/client-portal",
        "https://www.salesforce.com/client-portal",
    ),
    "client portal vs project management tool": _serp(
        "https://www.monday.com/compare",
        "https://www.asana.com/compare",
    ),
}
PEOPLE_ALSO_ASK = {
    "client project workspace": [
        "How much does a client portal cost?",
        "Can clients approve files online?",
    ],
    "client portal for agencies": ["Is a client portal secure?"],
}
KEYWORDS: dict[str, tuple[int, int | None]] = {
    # phrase: (monthly Bing impressions, difficulty). Free mode ignores this difficulty and
    # computes its own from the top 10 results (chapter 8).
    "client portal for agencies": (90, None),
    "client portal login": (60, None),
    "client project workspace": (40, None),
    "client approval software": (30, None),
    "file sharing with clients": (25, None),
    "project management software": (5000, None),
    "agency client tracking": (3, None),
}
TRANCO = {  # site popularity ranks from the Tranco list (1 = most popular site on the web)
    "reddit.com": 20,
    "microsoft.com": 10,
    "adobe.com": 50,
    "salesforce.com": 60,
    "hubspot.com": 300,
    "atlassian.com": 400,
    "asana.com": 700,
    "monday.com": 800,
    "zendesk.com": 900,
    "g2.com": 900,
    "clientflow.com": 150_000,
}

# ---------------------------------------------------------------------------------------
# 4. The pretend LLM: one small rule per prompt type (keyed by the Pydantic schema name)
# ---------------------------------------------------------------------------------------
SEEDS = [
    "client portal for agencies",
    "client project workspace",
    "file sharing with clients",
    "client approval software",
    "client portal login",
    "project management software",
    "agency client tracking",
]
# People searching "client portal login" want to sign in to a portal they already use.
NOT_A_FIT = {"client portal login"}


def _main_phrase(user: str) -> str:
    return re.search(r"MAIN PHRASE: (.+)", user).group(1).strip()  # type: ignore[union-attr]


def fit_verdicts(system: str, user: str) -> dict:
    phrases = [line[2:] for line in user.split("PHRASES:\n")[1].splitlines()]
    return {"verdicts": [{"phrase": p, "fits": p not in NOT_A_FIT} for p in phrases]}


def page_reading(system: str, user: str) -> dict:
    text = user.lower()
    topics = [t for t in PARA if t in text]
    if "client project" in text and not topics:  # our own page
        topics = ["client approvals", "shared task list", "project timeline"]
    kind = "listicle" if "best client portals" in text else "product"
    return {"page_type": kind, "topics": topics or ["client portals"], "questions": []}


def brief_draft(system: str, user: str) -> dict:
    main = _main_phrase(user)
    return {
        "titles": [
            f"{main.capitalize()} for Agencies | Emitii",
            f"{main.capitalize()}: Files, Approvals and Tasks | Emitii",
        ],
        "description": (
            f"A {main} where agencies share files, collect client approvals and track tasks. "
            "Invite clients in minutes, no email threads."
        ),
        "headings": [
            f"The {main} for agencies",
            "Share files with clients",
            "Client approvals in one place",
            "Pricing plans",
            "Is a client portal secure?",
        ],
    }


def _bullets(user: str, header: str) -> list[str]:
    """The "- item" lines under one HEADER: in the prompt the engine sent."""
    block = user.split(header + "\n")[1].split("\n\n")[0]
    return [line[2:] for line in block.splitlines() if line.startswith("- ") and line != "- none"]


def draft_out(system: str, user: str) -> dict:
    main = _main_phrase(user)
    topics = (_bullets(user, "MUST-COVER TOPICS:") + _bullets(user, "ALSO WORTH COVERING:"))[:4]
    return {
        "h1": f"The {main} for agencies",
        "intro": f"Emitii is a {main} where agencies share files, collect approvals and keep "
        "tasks visible to both sides.",
        "sections": [
            {"heading": t.capitalize(), "body": f"How Emitii handles {t}: [ADD: {t} details]"}
            for t in topics
        ],
        "faq": [
            {"question": q, "answer": "[ADD: your answer]"}
            for q in _bullets(user, "GAP QUESTIONS FOR THE FAQ:")
        ],
        "cta": "Start a free workspace for your next client project.",
    }


class CostedLLM(FakeLLM):
    """The fake LLM, but it reports a pretend cost for every call, like DeepSeekLLM does."""

    def __init__(self, handlers, cost_sink, usd_per_call: float = 0.0005) -> None:
        super().__init__(handlers)
        self.cost_sink = cost_sink
        self.usd_per_call = usd_per_call

    def structured(self, system, user, schema, tier="bulk"):
        self.cost_sink(self.usd_per_call, f"fake-{tier}")
        return super().structured(system, user, schema, tier)


def make_llm(cost_sink=lambda usd, label: None) -> CostedLLM:
    return CostedLLM(
        {
            "SeedPhrases": lambda s, u: {"phrases": SEEDS},
            "FitVerdicts": fit_verdicts,
            "PageReading": page_reading,
            "PassageLabels": fake_passage_labels,
            "NoiseTopics": lambda s, u: {"noise": []},
            "GapFits": all_gaps_relevant,
            "BriefDraft": brief_draft,
            "DraftOut": draft_out,
        },
        cost_sink,
    )


# ---------------------------------------------------------------------------------------
# 5. Put it together
# ---------------------------------------------------------------------------------------
def make_settings(**overrides) -> Settings:
    """Default settings, with the cache pointed at the scratch folder."""
    return Settings(cache_dir=SCRATCH / "cache", **overrides)


def make_run(**setting_overrides) -> Run:
    return Run(page_text=OUR_PAGE, settings=make_settings(**setting_overrides))


def make_deps(run: Run) -> Deps:
    pages = {u: _fetched(u, title, text) for u, (title, text) in COMPETITORS.items()}
    return Deps(
        search=FakeSearch(SERPS, paa=PEOPLE_ALSO_ASK),
        keywords=FakeKeywords(
            KEYWORDS,
            autocomplete={"client project workspace": ["client workspace app"]},
        ),
        fetcher=FakeFetcher(pages),
        llm=make_llm(run.add_cost),
        embed=FakeEmbed(),
        ranks=DictRanks(TRANCO),
        autocomplete=FakeAutocomplete(
            searched={"client project workspace", "client workspace app"},
            variants=["what is a client portal", "client portal vs project management tool"],
        ),
    )

# ========================================================================================
# 2. narrate(): prints what happened in a run, step by step (chapter 4)
# ========================================================================================
import time

from seo_engine.models import Run
from seo_engine.pipeline import PipelineDetails


def on_step(name: str, status: str, detail: str) -> None:
    """The pipeline calls this before and after each of its 6 steps (the web app uses it
    to move the progress bar)."""
    stamp = time.strftime("%H:%M:%S")
    print(f"  {stamp}  {name:<12} {status:<8} {detail}")


def heading(text: str) -> None:
    print(f"\n=== {text} " + "=" * max(0, 70 - len(text)))


def narrate(run: Run, details: PipelineDetails, llm_calls: list[str] | None = None) -> None:
    heading("Step 1: every search phrase considered, and why it was kept or dropped")
    for c in details.candidates:
        mark = "CHOSEN" if c.kept else "      "
        print(f"  {mark} {c.keyword:<42} from {'+'.join(c.sources):<13} {c.reason}")

    heading("Step 1: clusters, ranked by fit x demand x winnability")
    for cl in details.clusters:
        print(
            f"  {cl.head:<28} score {cl.score:.3f} = fit {cl.fit:.2f} x demand {cl.demand:.2f}"
            f" x win {cl.winnability:.2f}   members: {cl.members}"
        )

    heading("Step 2: Google results per chosen phrase")
    for serp in details.serps:
        print(f"  '{serp.phrase}': {[i.domain for i in serp.items]}")
        if serp.people_also_ask:
            print(f"      People Also Ask: {serp.people_also_ask}")

    heading("Step 3: competitor pages kept and dropped")
    for p in run.competitors:
        print(f"  KEPT     {p.source:<10} {p.url:<45} {p.page_type}, topics {p.topics}")
    for d in details.dropped:
        print(f"  DROPPED  google#{d.rank:<3} {d.url:<45} {d.reason}")
    print(f"  Our page reads as: {details.ours_page_type}. Intent verdict: {details.intent.verdict}")

    heading("Step 4: topic coverage (who covers what)")
    print(f"  {'topic':<26}{'bucket':<8}{'competitors':<13}{'our passages'}")
    for t in run.coverage:
        print(f"  {t.topic:<26}{t.bucket:<8}{f'{t.covered_by} of {t.total}':<13}{t.ours_passages}")
    print("\n  Gaps (few competitors answer them, searchers ask them):")
    for g in run.gaps:
        print(f"   - {g.topic}   [{g.evidence}]")

    brief = run.brief
    heading("Step 4: the score, with its arithmetic")
    print(f"  {brief.score}/100\n  {brief.score_arithmetic}")

    heading("Step 5: the brief (titles and description from the LLM, the rest from code)")
    for t, s in zip(brief.titles, details.snippets, strict=False):
        print(f"  title: {t!r}  ({s.title_px}px, {'ok' if s.title_ok else 'too long'})")
    print(f"  description: {brief.description!r}")
    print(f"  headings: {brief.headings}")
    print(f"  checklist: {brief.checklist}")

    heading("Step 6: the suggested draft")
    draft = brief.draft
    print(f"  H1: {draft.h1}\n  {draft.word_count} words, placeholders to fill: {draft.placeholders}")
    for check in draft.checks:
        print(f"   {'PASS' if check.ok else 'FAIL'}  {check.label}  {check.detail}")

    heading("Cost log (every paid call adds a line; these are pretend prices offline)")
    by_label: dict[str, int] = {}
    for entry in run.costs:
        by_label[entry.label] = by_label.get(entry.label, 0) + 1
    print(f"  {len(run.costs)} LLM calls: {by_label}   total ${run.cost_usd:.4f}")
    if llm_calls is not None:
        print(f"  LLM calls by prompt type: {dict((n, llm_calls.count(n)) for n in llm_calls)}")

# ========================================================================================
# 3. The exercises, one function per exercise (exNN = chapter NN)
# ========================================================================================
def ex02_embeddings() -> None:
    """Chapter 2: embeddings and cosine similarity, from tiny hand-made vectors. Free, offline.

        .venv/bin/python docs/guide_exercises.py ex02_embeddings

    An embedding is a list of numbers that stands for the meaning of a text. Texts with similar
    meaning get vectors pointing in similar directions. Cosine similarity measures that: 1 means
    the same direction, 0 means unrelated.

    Things to try:
    - Add your own 3-number vectors to `toy` and predict the similarity before running.
    - Add more phrases to `texts` below. FakeEmbed only sees shared words; a real embedding model
      (Gemini) also knows that "customer" and "client" mean nearly the same thing.
    """

    import math


    from seo_engine.providers.embeddings import cosine, normalise

    print("--- 1. Tiny vectors. Pretend the 3 numbers mean (about clients, about files, about food).")
    toy = {
        "client portal": [0.9, 0.3, 0.0],
        "customer portal": [0.8, 0.4, 0.0],
        "file sharing": [0.2, 0.9, 0.0],
        "banana bread": [0.0, 0.0, 1.0],
    }
    unit = {k: normalise(v) for k, v in toy.items()}
    for k, v in unit.items():
        length = math.sqrt(sum(x * x for x in v))
        print(f"  {k:<16} normalised {[round(x, 3) for x in v]}  length {length:.3f}")

    print("\n--- 2. Cosine similarity to 'client portal' (unit vectors, so it is just a dot product).")
    for k in unit:
        print(f"  {k:<16} {cosine(unit['client portal'], unit[k]):.3f}")

    print("\n--- 3. FakeEmbed, the test stand-in for Gemini: similar when texts share words.")
    emb = FakeEmbed()
    texts = ["client portal", "client portals", "file sharing with clients", "banana bread recipe"]
    vecs = emb.embed(texts)
    for t, v in zip(texts, vecs, strict=True):
        print(f"  {t:<28} {cosine(vecs[0], v):.3f}")
    print(f"  (vectors have {len(vecs[0])} numbers each; Gemini's have 768)")


def ex02_llm_json_offline() -> None:
    """Chapter 2: how one LLM "JSON mode" call works, with a fake DeepSeek server. Free, offline.

        .venv/bin/python docs/guide_exercises.py ex02_llm_json_offline

    The real DeepSeekLLM class (src/seo_engine/providers/llm.py) is used unchanged. Only its
    httpx client is swapped for one with a MockTransport: a function that plays the server.
    The fake server answers badly the first time and correctly the second time, so you can see
    the "retry once, then raise" rule and the cost log.

    Things to try:
    - Make the server answer badly twice (return BAD every time) and read the LLMOutputError.
    - Change "completion_tokens" to 5000 and watch the cost change.
    """

    import json

    import httpx
    from pydantic import BaseModel

    from seo_engine.config import ModelSettings
    from seo_engine.models import Run
    from seo_engine.providers.llm import DeepSeekLLM


    class Topics(BaseModel):
        topics: list[str]


    BAD = "Sure! Here are the topics: file sharing, approvals"  # not JSON at all
    GOOD = '{"topics": ["file sharing", "client approvals"]}'
    replies = [BAD, GOOD]


    def fake_server(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        print(f"  server got a request: model={body['model']}, {len(body['messages'])} messages,")
        print(f"    response_format={body.get('response_format')}")
        print(f"    last message ({body['messages'][-1]['role']}): "
              f"{body['messages'][-1]['content'][:70]!r}")
        content = replies.pop(0)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"role": "assistant", "content": content}}],
                "usage": {"prompt_cache_miss_tokens": 1000, "prompt_cache_hit_tokens": 0,
                          "completion_tokens": 500},
            },
        )


    run = Run(page_text="x")
    client = httpx.Client(base_url="https://api.deepseek.com", transport=httpx.MockTransport(fake_server))
    llm = DeepSeekLLM(ModelSettings(), api_key="fake-key", cost_sink=run.add_cost, client=client)

    print("--- Asking for Topics (the system prompt ends with the JSON schema of Topics)")
    out = llm.structured("List the subtopics of the page.", "Share files and approvals.", Topics)
    print("\n--- Result, already a validated Topics object:")
    print(" ", out)
    print("\n--- Cost log: one entry per call, even the failed one")
    for c in run.costs:
        print(f"  {c.label}: ${c.usd:.6f}")
    print(f"  total: ${run.cost_usd:.6f}   (1000 in x $0.28/M + 500 out x $0.42/M = $0.00049 per call)")


def ex02_protocol_and_fakes() -> None:
    """Chapter 2: Protocols ("duck typing with a contract") and why fakes work. Free, offline.

        .venv/bin/python docs/guide_exercises.py ex02_protocol_and_fakes

    The serp_top tool wants a `SearchProvider`: anything with a method
    `top(phrase, country, n) -> SerpResults`. It never checks the class. So a ten-line class you
    write here works exactly like the real Serper or Gemini provider.

    Things to try:
    - Make MySearch return 7 listicles and 3 products, and pass our_type="product".
    - Rename `top` to `search` and run it: the program crashes with AttributeError, because
      Protocols are only checked by type-checking tools, not at run time.
    """

    from seo_engine.providers.search import SerpItem, SerpResults
    from seo_engine.tools.serp_top import serp_top


    class MySearch:
        """Not a subclass of anything. It just has the right method."""

        def top(self, phrase: str, country: str, n: int) -> SerpResults:
            kinds = ["listicle"] * 4 + ["product"] * 6
            items = [
                SerpItem(rank=i + 1, url=f"https://site{i}.com/", domain=f"site{i}.com", page_type=k)
                for i, k in enumerate(kinds[:n])
            ]
            return SerpResults(phrase=phrase, country=country, items=items)


    result = serp_top(MySearch(), "client portal", "US", n=10, our_type="product")
    print("results:", [i.page_type for i in result.serp.items])
    print("intent :", result.intent)

    result = serp_top(MySearch(), "client portal", "US", n=10, our_type="guide")
    print("\nIf our page were a guide:")
    print("flag   :", result.intent.flag)


def ex02_pydantic() -> None:
    """Chapter 2: Pydantic models, using the engine's real models. Free, offline.

        .venv/bin/python docs/guide_exercises.py ex02_pydantic

    Shows: building a model, defaults, a ValidationError, dumping to a dict and JSON, reading
    JSON back, model_copy, and the JSON schema an LLM is shown.

    Things to try:
    - Give the Brief a score of 101 and read the error.
    - Remove `intent=` from the Phrase and see which field Pydantic complains about.
    - Print Phrase.model_json_schema() in full.
    """

    import json

    from pydantic import ValidationError

    from seo_engine.models import Brief, Phrase, Run

    print("--- 1. Build a model. Missing optional fields get their defaults.")
    p = Phrase(text="client project workspace", volume=40, difficulty=13, intent="commercial")
    print(p)
    print("volume_source default:", p.volume_source, "| cluster default:", p.cluster)

    print("\n--- 2. Pydantic converts obvious types for you: the string '40' becomes the int 40.")
    print(Phrase(text="x", volume="40", difficulty=1, intent="?").volume + 1)

    print("\n--- 3. Bad data raises ValidationError with every problem listed.")
    try:
        Brief(
            phrases=[],  # min_length=1 in models.py
            titles=[],
            description="",
            must_cover=[],
            gaps=[],
            headings=[],
            intent_flag=None,
            score=120,  # le=100 in models.py
            checklist={},
        )
    except ValidationError as exc:
        print(exc)

    print("--- 4. model_dump() gives a dict, model_dump_json() gives a JSON string.")
    print(p.model_dump())
    text = p.model_dump_json()
    print(text)

    print("\n--- 5. model_validate_json() reads JSON back into a checked object.")
    again = Phrase.model_validate_json(text)
    print("same as before:", again == p)

    print("\n--- 6. model_copy(update=...) makes a changed copy; the original is untouched.")
    harder = p.model_copy(update={"difficulty": 55})
    print("copy:", harder.difficulty, "| original:", p.difficulty)

    print("\n--- 7. A Run starts nearly empty and gets filled in by the pipeline.")
    run = Run(page_text="Emitii is a client project workspace.")
    print("country:", run.settings.country, "| phrases:", run.phrases, "| cost:", run.cost_usd)
    run.add_cost(0.002, "deepseek-chat")
    run.add_cost(0.0006, "deepseek-chat")
    print("after two add_cost calls:", run.cost_usd, run.costs)

    print("\n--- 8. The JSON schema of a model (this is what the LLM is shown, chapter 7).")
    print(json.dumps(Phrase.model_json_schema())[:300], "...")


def ex03_settings_tour() -> None:
    """Chapter 3: a tour of every setting and threshold, with their default values. Free, offline.

        .venv/bin/python docs/guide_exercises.py ex03_settings_tour

    Every depth, cost and behaviour knob lives in src/seo_engine/config.py (CLAUDE.md rule 9).
    This prints the whole tree so you can see them all in one place, and shows which API keys
    your .env provides (true/false only, never the keys themselves).

    Things to try:
    - Run with a different site strength: change `Settings()` to Settings(site_strength="growing")
      and compare the printed difficulty ceiling.
    - Find one threshold below in config.py and in the tool that uses it (grep for its name).
    """

    from seo_engine.config import PROJECT_ROOT, Secrets, Settings

    s = Settings()
    print("PROJECT_ROOT:", PROJECT_ROOT)

    print("\n--- Run settings (one Settings object per run)")
    for name, value in s.model_dump(exclude={"models", "thresholds"}).items():
        print(f"  {name:<22} {value}")

    print("\n--- Model settings (which LLM and embedding model each tier uses)")
    for name, value in s.models.model_dump().items():
        print(f"  {name:<28} {value}")

    print("\n--- Thresholds (the algorithm's tunable numbers)")
    for name, value in s.thresholds.model_dump().items():
        text = str(value)
        print(f"  {name:<30} {text if len(text) < 60 else text[:57] + '...'}")

    print(f"\nDifficulty ceiling for a '{s.site_strength}' site:",
          s.thresholds.difficulty_ceiling[s.site_strength])

    print("\n--- Keys found in .env (booleans only)")
    sec = Secrets()
    for name in ["deepseek_api_key", "gemini_api_key", "serper_api_key", "bing_webmaster_api_key",
                 "dataforseo_password"]:
        print(f"  {name:<24} {bool(getattr(sec, name).get_secret_value())}")
    print("  repr of a secret prints stars, not the key:", repr(sec.deepseek_api_key)[:40])


def ex04_full_run_offline() -> None:
    """Chapter 4: run the whole engine, start to finish, on the offline world. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex04_full_run_offline

    It calls the real `run_pipeline` (src/seo_engine/pipeline.py), the same function the web app
    uses. Only the providers are fake (the offline world (top of this file)). After it prints, open
    cache/guide_scratch/run.json to see the full `Run` object the engine filled in.

    Things to try afterwards:
    - Change site_strength="new" to "growing" below and run again. Why does nothing change?
    - Change phrases_per_run to 1. Why are fewer competitor pages kept?
    - In the offline world (top of this file), empty the NOT_A_FIT set. What gets chosen now?
    """

    from seo_engine.pipeline import run_pipeline

    run = make_run(site_strength="new", phrases_per_run=3)
    deps = make_deps(run)

    heading("Running the pipeline")
    details = run_pipeline(run, deps, on_step, today=TODAY)
    narrate(run, details, deps.llm.calls)

    SCRATCH.mkdir(parents=True, exist_ok=True)
    (SCRATCH / "run.json").write_text(run.model_dump_json(indent=2))
    print(f"\nSaved the full Run object to {SCRATCH / 'run.json'}")


def ex04_live_run() -> None:
    """Chapter 4 (optional): the same narrated run, but with the REAL providers.

        .venv/bin/python docs/guide_exercises.py ex04_live_run

    This one is live. It reads your keys from .env, searches Google (via Serper or Gemini),
    fetches real competitor pages and calls DeepSeek. Past runs cost between $0.018 and $0.041
    and took about 1.5 to 3 minutes. It asks before spending anything.

    It builds providers exactly the way the API does, with `from_env` in src/seo_engine/deps.py.
    """

    from seo_engine.deps import from_env
    from seo_engine.models import Run
    from seo_engine.pipeline import run_pipeline

    answer = input("This spends a few cents of DeepSeek credit. Type yes to continue: ")
    if answer.strip().lower() != "yes":
        raise SystemExit("Stopped. Nothing was spent.")

    run = Run(page_text=OUR_PAGE)  # default Settings, same as the web form's defaults
    deps = from_env(run)

    heading("Running the pipeline (live)")
    details = run_pipeline(run, deps, on_step)
    narrate(run, details)

    SCRATCH.mkdir(parents=True, exist_ok=True)
    (SCRATCH / "live_run.json").write_text(run.model_dump_json(indent=2))
    print(f"\nSaved the full Run object to {SCRATCH / 'live_run.json'}")


def ex05_run_and_costs() -> None:
    """Chapter 5: the Run object, settings, cost logging and validation. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex05_run_and_costs

    What it shows:
    1. A Run starts almost empty; its Settings come with defaults you can override.
    2. Thresholds sit inside Settings, and so do the model names.
    3. add_cost() is how every paid call reports its price (CLAUDE.md rule 5).
    4. A Run turns into JSON and back without losing anything (the API stores runs this way).
    5. Pydantic refuses a Brief that breaks its rules (score 120, no phrases).

    Things to try:
    - Override a threshold: Settings(thresholds=Thresholds(min_bing_impressions=50)).
    - Set site_strength="huge" and read the error Pydantic gives you.
    - Call add_cost from 4 threads at once and check the total (see the lock in models.py).
    """

    import threading

    from pydantic import ValidationError

    from seo_engine.config import Settings, Thresholds
    from seo_engine.models import Brief, Phrase, Run

    print("== 1. A new Run ==")
    run = Run(page_text="Emitii is a client project workspace for agencies.")
    print("phrases:", run.phrases, "| brief:", run.brief, "| cost:", run.cost_usd)
    print("country:", run.settings.country, "| site_strength:", run.settings.site_strength,
          "| data_mode:", run.settings.data_mode)

    print("\n== 2. Overriding settings ==")
    s = Settings(country="GB", site_strength="established")
    ceiling = s.thresholds.difficulty_ceiling[s.site_strength]
    print(f"GB, established site -> difficulty ceiling {ceiling}")
    strict = Settings(thresholds=Thresholds(min_bing_impressions=50))
    print("custom Bing floor:", strict.thresholds.min_bing_impressions,
          "| default floor:", Settings().thresholds.min_bing_impressions)
    print("LLM for judgment calls:", s.models.judgment, "| for bulk calls:", s.models.bulk)

    print("\n== 3. Logging costs ==")
    run.add_cost(0.0021, "deepseek-reasoner")
    run.add_cost(0.0004, "deepseek-chat")
    run.add_cost(0.0, "gemini.grounding.gemini-2.5-flash")  # free tier still gets a line
    for entry in run.costs:
        print(f"  {entry.label:<36} ${entry.usd}")
    print("total:", run.cost_usd)

    threads = [threading.Thread(target=lambda: [run.add_cost(0.001, "t") for _ in range(250)])
               for _ in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    print(f"after 4 threads x 250 calls of $0.001: {len(run.costs)} entries, total {run.cost_usd}")

    print("\n== 4. JSON round trip ==")
    SCRATCH.mkdir(parents=True, exist_ok=True)
    path = SCRATCH / "ex05_run.json"
    path.write_text(run.model_dump_json(indent=2))
    again = Run.model_validate_json(path.read_text())
    print(f"wrote {path.name} ({path.stat().st_size} bytes); read back equal: {again == run}")

    print("\n== 5. Validation errors ==")
    phrase = Phrase(text="client project workspace", volume=40, difficulty=13, intent="commercial")
    good = dict(phrases=[phrase], titles=["t"], description="d", must_cover=[], gaps=[],
                headings=["h"], intent_flag=None, score=55, checklist={})
    print("valid brief score:", Brief(**good).score)
    for bad in ({"score": 120}, {"phrases": []}, {"phrases": [phrase] * 4}):
        try:
            Brief(**{**good, **bad})
        except ValidationError as exc:
            err = exc.errors()[0]
            print(f"  {bad!s:<40.40} -> {err['loc'][0]}: {err['msg']}")
    try:
        Settings(site_strength="huge")
    except ValidationError as exc:
        print("  site_strength='huge' ->", exc.errors()[0]["msg"])


def ex06_cache_and_retry() -> None:
    """Chapter 6: the daily cache and the retrying HTTP call. Free, no network.

        .venv/bin/python docs/guide_exercises.py ex06_cache_and_retry

    What it shows:
    1. DailyCache writes one JSON file per (namespace, inputs, day). You see the exact path.
    2. The same inputs in a different key order hit the same file (sort_keys=True).
    3. A new day means a new folder, so yesterday's answers are never reused.
    4. request_with_retry against a fake server that fails twice with 503, then answers.
       The real code sleeps 1 s, then 2 s; here `sleep` is replaced by a print.
    5. A 404 is not retried: it fails at once.

    Things to try:
    - Change the fake server to return 429 four times. What happens on the last attempt?
    - Pass attempts=2 to request_with_retry and watch it give up sooner.
    - Look inside cache/guide_scratch/ex06_cache/ after running.
    """

    import shutil
    from datetime import date

    import httpx

    from seo_engine.providers.base import DailyCache, request_with_retry

    root = SCRATCH / "ex06_cache"
    shutil.rmtree(root, ignore_errors=True)

    print("== 1. One cache entry ==")
    cache = DailyCache(root, today=lambda: date(2026, 9, 25))
    key = {"q": "client portal", "gl": "us", "num": 10}
    print("before set:", cache.get("serper", key))
    cache.set("serper", key, {"organic": [{"link": "https://portalpro.io/"}]})
    for path in root.rglob("*.json"):
        print("file:", path.relative_to(SCRATCH))
    print("after set:", cache.get("serper", key))

    print("\n== 2. Key order does not matter ==")
    same_key_other_order = {"num": 10, "gl": "us", "q": "client portal"}
    print("hit:", cache.get("serper", same_key_other_order) is not None)
    print("num 20 instead of 10 is a different entry:",
          cache.get("serper", {**key, "num": 20}))

    print("\n== 3. Tomorrow starts empty ==")
    tomorrow = DailyCache(root, today=lambda: date(2026, 9, 26))
    print("tomorrow's get:", tomorrow.get("serper", key))

    print("\n== 4. Retrying a flaky server ==")
    replies = iter([503, 503, 200])


    def flaky(request: httpx.Request) -> httpx.Response:
        status = next(replies)
        print(f"  server got {request.method} {request.url.path}, answers {status}")
        return httpx.Response(status, json={"ok": status == 200})


    client = httpx.Client(transport=httpx.MockTransport(flaky), base_url="https://api.example.com")
    resp = request_with_retry(
        client, "POST", "/search", json={"q": "x"}, sleep=lambda s: print(f"  (would sleep {s} s)")
    )
    print("final:", resp.status_code, resp.json())

    print("\n== 5. A 404 is not retried ==")
    calls = []


    def missing(request: httpx.Request) -> httpx.Response:
        calls.append(1)
        return httpx.Response(404)


    client = httpx.Client(transport=httpx.MockTransport(missing), base_url="https://api.example.com")
    try:
        request_with_retry(client, "GET", "/nope", sleep=lambda s: None)
    except httpx.HTTPStatusError as exc:
        print(f"raised {type(exc).__name__} after {len(calls)} call(s): {exc.response.status_code}")


def ex06_fallback_search() -> None:
    """Chapter 6: search providers, the Protocol idea and the free-mode fallback. Free, no network.

        .venv/bin/python docs/guide_exercises.py ex06_fallback_search

    What it shows:
    1. parse_serper turns a raw Serper.dev reply into the engine's own SerpResults shape,
       guessing a page type for each result from its URL and title.
    2. SerperSearch with no API key raises SearchUnavailable, and FallbackSearch moves on.
    3. SerperSearch talking to a fake Serper server (httpx.MockTransport): the second identical
       call comes from the daily cache, so the server is called once.
    4. When every provider is unavailable, FallbackSearch lists all the reasons.

    Things to try:
    - In section 3, make the fake server answer 403 ("no credits"). Which provider answers now?
    - Add "images": [{}] to RAW and see which feature appears.
    - Write your own provider class with a top() method and put it first in the list.
    """

    import shutil
    from datetime import date

    import httpx

    from seo_engine.providers.base import DailyCache
    from seo_engine.providers.search import FallbackSearch, SearchUnavailable, SerpItem, SerpResults
    from seo_engine.providers.serper import SERPER_URL, SerperSearch, parse_serper

    RAW = {  # the shape Serper.dev documents for POST /search
        "organic": [
            {"title": "PortalPro | Client Workspace Software", "link": "https://portalpro.io/",
             "position": 1, "snippet": "One workspace per client."},
            {"title": "12 Best Client Portals (2026)", "link": "https://reviewhub.com/best-portals",
             "position": 2},
            {"title": "How to set up a client portal", "link": "https://blog.io/how-to-client-portal",
             "position": 3},
            {"title": "r/agency", "link": "https://www.reddit.com/r/agency/1", "position": 4},
        ],
        "peopleAlsoAsk": [{"question": "Is a client portal secure?"}],
        "relatedSearches": [{"query": "client portal free"}],
        "videos": [{"title": "demo"}],
    }

    print("== 1. parse_serper ==")
    serp = parse_serper(RAW, "client portal", "US", 10)
    for item in serp.items:
        print(f"  #{item.rank} {item.domain:<16} {item.page_type:<10} {item.title}")
    print("  features:", serp.features)
    print("  PAA:", serp.people_also_ask, "| related:", serp.related_searches)


    class HandMadeSearch:
        """Any class with this top() method counts as a SearchProvider (a Protocol, chapter 2)."""

        def top(self, phrase: str, country: str, n: int) -> SerpResults:
            items = [SerpItem(rank=1, url="https://example.com/", domain="example.com")]
            return SerpResults(phrase=phrase, country=country, source="hand-made", items=items)


    cache = DailyCache(SCRATCH / "ex06_search_cache", today=lambda: date(2026, 9, 25))
    shutil.rmtree(cache.root, ignore_errors=True)

    print("\n== 2. No Serper key: fall back ==")
    search = FallbackSearch([SerperSearch(api_key="", cache=cache), HandMadeSearch()])
    result = search.top("client portal", "US", 10)
    print("answered by:", result.source, [i.url for i in result.items])

    print("\n== 3. A fake Serper server, and the cache ==")
    calls = []


    def fake_serper(request: httpx.Request) -> httpx.Response:
        calls.append(request.url.path)
        return httpx.Response(200, json=RAW)


    client = httpx.Client(transport=httpx.MockTransport(fake_serper), base_url=SERPER_URL)
    serper = SerperSearch(api_key="pretend-key", cache=cache, client=client)
    first = serper.top("client portal", "US", 10)
    second = serper.top("client portal", "US", 10)
    twenty = serper.top("client portal", "US", 20)  # "num": 20 is a different request
    print(f"3 calls to top(), {len(calls)} reached the server; results {len(first.items)}, "
          f"{len(second.items)}, {len(twenty.items)}")

    print("\n== 4. Nothing available ==")
    try:
        FallbackSearch([SerperSearch("", cache), SerperSearch("", cache)]).top("x", "US", 10)
    except SearchUnavailable as exc:
        print("SearchUnavailable:", exc)


def ex07_bing_and_autocomplete() -> None:
    """Chapter 7: free-mode demand data (Bing, Google autocomplete) and site strength (Tranco).
    Free, no network: every "server" here is an httpx.MockTransport.

        .venv/bin/python docs/guide_exercises.py ex07_bing_and_autocomplete

    What it shows:
    1. monthly_from_weekly: Bing gives weekly impressions, one dated row per week, and leaves
       out weeks with none. The engine adds up the last 12 weeks by date, divides by 12 and
       multiplies by 52/12 (changed 2026-09-29, chapter 20: it used to average the rows).
    2. BingKeywords.metrics and .suggestions against a fake Bing API, and what happens with
       no API key (every volume is 0).
    3. GoogleAutocomplete: is_searched() ("does the phrase autocomplete to itself?") and
       variants() (question patterns like "what is <seed>").
    4. Tranco lookups: a subdomain falls back to its parent domain.

    Things to try:
    - Give monthly_from_weekly 20 dated rows. Which 12 weeks does it use?
    - Add "client portal" to the fake suggest table for "client portal software" and check
      is_searched("client portal software").
    - Look up "docs.github.com" in the DictRanks below.
    """

    import shutil
    from datetime import date, timedelta

    import httpx

    from seo_engine.providers.autocomplete import GoogleAutocomplete
    from seo_engine.providers.base import DailyCache
    from seo_engine.providers.bing import BING_URL, BingKeywords, monthly_from_weekly
    from seo_engine.providers.tranco import DictRanks, candidates

    today = date(2026, 9, 25)
    cache = DailyCache(SCRATCH / "ex07_demand_cache", today=lambda: today)
    shutil.rmtree(cache.root, ignore_errors=True)

    def dated(counts: list[int], last: date = date(2026, 9, 19)) -> list[dict]:
        """Rows shaped like the live API: one per week, oldest first, WCF dates."""
        epoch = date(1970, 1, 1)
        n = len(counts)
        return [
            {"Impressions": c, "Date": f"/Date({(last - timedelta(weeks=n - 1 - i) - epoch).days * 86_400_000})/"}
            for i, c in enumerate(counts)
        ]

    print("== 1. Weekly to monthly ==")
    weeks = dated([0, 0, 5, 8, 10, 12, 9, 11, 10, 14, 12, 13, 15, 9])
    recent = [w["Impressions"] for w in weeks][-12:]
    total = sum(recent)
    print("last 12 of", len(weeks), "weeks:", recent)
    print(f"total {total} / 12 weeks x 52/12 = {total / 12 * 52 / 12:.1f} -> {monthly_from_weekly(weeks, today)}")
    sparse = dated([1])  # live, 28 Sep 2026: "client portal software" came back as ONE row
    print("one row of 1 impression (Bing leaves out empty weeks):",
          monthly_from_weekly(sparse, today), "a month (averaging the rows said 4)")

    print("\n== 2. Bing, with a fake API ==")


    def fake_bing(request: httpx.Request) -> httpx.Response:
        q = request.url.params["q"]
        if request.url.path.endswith("GetKeywordStats"):
            per_week = {"client portal": 30, "client project workspace": 4}.get(q, 0)
            return httpx.Response(200, json={"d": dated([per_week] * 12)})
        rows = [{"Query": "Client Portal Software", "Impressions": 300},
                {"Query": "client portal app", "Impressions": 60}]
        return httpx.Response(200, json={"d": rows})


    bing_http = httpx.Client(transport=httpx.MockTransport(fake_bing), base_url=BING_URL)
    autocomplete_stub = GoogleAutocomplete(cache)  # BingKeywords.autocomplete() delegates to it
    bing = BingKeywords("pretend-key", cache, autocomplete_stub, client=bing_http,
                        today=date(2026, 9, 25))
    for m in bing.metrics(["client portal", "Client  Project Workspace", "purple llamas"], "US"):
        print(f"  {m.keyword:<26} volume {m.volume:>4}  difficulty {m.difficulty}")
    print("  related (3 months of impressions / 3):",
          [(r.keyword, r.volume) for r in bing.suggestions("client portal", "US")])
    no_key = BingKeywords("", cache, autocomplete_stub)
    print("  without a key:", [(m.keyword, m.volume) for m in no_key.metrics(["client portal"], "US")])

    print("\n== 3. Google autocomplete, with a fake suggest endpoint ==")
    TABLE = {
        "client portal": ["client portal", "client portal login", "client portal software"],
        "what is client portal": ["what is a client portal"],
        "client portal vs": ["client portal vs extranet", "client portal vs crm"],
    }


    def fake_suggest(request: httpx.Request) -> httpx.Response:
        q = request.url.params["q"]
        return httpx.Response(200, json=[q, TABLE.get(q, []), [], {}])


    suggest_http = httpx.Client(transport=httpx.MockTransport(fake_suggest))
    ac = GoogleAutocomplete(cache, client=suggest_http)
    print("  suggest('Client  Portal'):", ac.suggest("Client  Portal", "US"))
    print("  is_searched('client portal'):", ac.is_searched("client portal", "US"))
    print("  is_searched('purple client portal'):", ac.is_searched("purple client portal", "US"))
    print("  variants:", ac.variants("client portal", "US", ["what is", "how to"], ["vs"]))

    print("\n== 4. Tranco rank lookups ==")
    ranks = DictRanks({"hubspot.com": 300, "example.co.uk": 45_000})
    for domain in ["blog.hubspot.com", "www.example.co.uk", "tiny-agency.io"]:
        print(f"  {domain:<20} tries {candidates(domain)} -> rank {ranks.rank(domain)}")


def ex07_fetcher_offline() -> None:
    """Chapter 7: the page fetcher, against a fake website. Free, no network.

        .venv/bin/python docs/guide_exercises.py ex07_fetcher_offline

    What it shows:
    1. A normal article page: robots.txt allows it, trafilatura cleans it, headings and the
       schema.org type come out.
    2. A page robots.txt forbids: the fetcher never asks for it.
    3. A JavaScript "shell" page with almost no text: status too_short without a browser,
       and "ok via headless" when a (fake) browser renders it.
    4. visible_text() on a landing page: keeps the copy, drops menus, cookie banners, footers.
    5. What was cached, and what was not (too_short is retried next time).

    The fake website serves the HTML files the unit tests use (tests/fixtures/html/).

    Things to try:
    - Change the robots.txt for shop.com to "User-agent: *\nDisallow: /private".
    - Make the fake server answer 403 for robots.txt. What does allowed() decide?
    - Raise Thresholds.min_clean_words to 400 and fetch the article again.
    """

    import json
    import shutil
    from datetime import date

    import httpx

    from seo_engine.config import Settings
    from seo_engine.providers.base import DailyCache
    from seo_engine.providers.fetcher import HttpFetcher, visible_text

    HTML = REPO / "tests" / "fixtures" / "html"
    ARTICLE = (HTML / "article.html").read_text()
    JS_SHELL = (HTML / "js_shell.html").read_text()

    SITE = {  # url -> (status, content type, body)
        "https://moxo.com/robots.txt": (404, "text/plain", ""),
        "https://moxo.com/blog/client-portal": (200, "text/html", ARTICLE),
        "https://shop.com/robots.txt": (200, "text/plain", "User-agent: *\nDisallow: /"),
        "https://shop.com/secret": (200, "text/html", ARTICLE),
        "https://app.com/robots.txt": (404, "text/plain", ""),
        "https://app.com/": (200, "text/html", JS_SHELL),
        "https://files.com/robots.txt": (404, "text/plain", ""),
        "https://files.com/guide.pdf": (200, "application/pdf", "%PDF-1.7"),
    }
    requested: list[str] = []


    def fake_web(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        requested.append(url)
        status, kind, body = SITE.get(url, (404, "text/plain", "not found"))
        return httpx.Response(status, text=body, headers={"content-type": kind})


    settings = Settings(cache_dir=SCRATCH / "ex07_cache")
    shutil.rmtree(settings.cache_dir, ignore_errors=True)
    cache = DailyCache(settings.cache_dir, today=lambda: date(2026, 9, 25))


    def fetcher(render=None) -> HttpFetcher:
        client = httpx.Client(transport=httpx.MockTransport(fake_web), follow_redirects=True)
        return HttpFetcher(settings, cache, client=client, render=render)


    def show(page) -> None:
        print(f"  status={page.status} method={page.method} words={page.word_count} "
              f"schema={page.schema_type!r}")
        print(f"  title={page.title!r}")
        print(f"  headings={page.headings[:4]}")
        if page.reason:
            print(f"  reason={page.reason!r}")


    print("== 1. An article ==")
    show(fetcher().fetch("https://moxo.com/blog/client-portal"))

    print("\n== 2. Blocked by robots.txt ==")
    show(fetcher().fetch("https://shop.com/secret"))
    print("  was /secret itself requested?", "https://shop.com/secret" in requested)

    print("\n== 3. A JavaScript shell page ==")
    print("  without a browser:")
    show(fetcher(render=None).fetch("https://app.com/"))
    print("  with a (fake) headless browser that returns the rendered HTML:")
    show(fetcher(render=lambda url, ua, timeout: ARTICLE).fetch("https://app.com/"))

    print("\n== 4. Not a web page ==")
    show(fetcher().fetch("https://files.com/guide.pdf"))

    print("\n== 5. visible_text on a landing page ==")
    LANDING = """<html><body>
  <nav><a>Home</a><a>Pricing</a></nav>
  <div id="cookie-consent"><p>We use cookies to improve your experience.</p></div>
  <section><h1>One workspace for client projects</h1>
    <p>Share files and approvals with every client.</p>
    <ul><li>Tasks both sides can see</li><li>Progress without meetings</li></ul></section>
  <form><input placeholder="Email"><button>Subscribe</button></form>
  <footer><p>© 2026 Emitii</p></footer>
</body></html>"""
    text, headings = visible_text(LANDING)
    print("  lines:", text.splitlines())
    print("  headings:", headings)

    print("\n== 6. What went into the cache ==")
    for path in sorted(settings.cache_dir.rglob("*.json")):
        saved = json.loads(path.read_text())
        print(f"  {path.parent.name}/{path.name[:8]}...  {saved['status']:<15} {saved['url']}")
    print("  (the too_short answer for app.com was not saved, so the second try fetched again)")


def ex07_llm_offline() -> None:
    """Chapter 7: the LLM and embedding providers, against fake servers. Free, no network.

        .venv/bin/python docs/guide_exercises.py ex07_llm_offline

    What it shows:
    1. The exact request DeepSeekLLM sends: model, JSON mode, and the JSON schema appended to
       the system prompt. Also: no temperature for deepseek-reasoner.
    2. The cost formula: cache-miss tokens, cache-hit tokens and output tokens, each at its price.
    3. Bad JSON once: the provider tells the model what was wrong and asks again. Bad twice:
       LLMOutputError (CLAUDE.md rule 2: retry once, then raise).
    4. GeminiEmbeddings: one batch call, vectors scaled to length 1, duplicates sent once,
       and a second call for the same text served from the daily cache.

    Things to try:
    - Change the fake usage numbers and check the cost by hand.
    - Make the fake DeepSeek always answer "{}" and read the LLMOutputError message.
    - Call embed() with 150 different texts: how many batches go to the fake server?
    """

    import json
    import math
    import shutil
    from datetime import date

    import httpx
    from pydantic import BaseModel

    from seo_engine.config import ModelSettings
    from seo_engine.models import Run
    from seo_engine.providers.base import DailyCache
    from seo_engine.providers.embeddings import GEMINI_URL, GeminiEmbeddings, cosine
    from seo_engine.providers.llm import DeepSeekLLM, LLMOutputError


    class Topics(BaseModel):
        topics: list[str]


    replies: list[str] = []
    sent: list[dict] = []


    def fake_deepseek(request: httpx.Request) -> httpx.Response:
        sent.append(json.loads(request.content))
        usage = {"prompt_tokens": 1000, "prompt_cache_hit_tokens": 200,
                 "prompt_cache_miss_tokens": 800, "completion_tokens": 500}
        msg = {"role": "assistant", "content": replies.pop(0)}
        return httpx.Response(200, json={"choices": [{"message": msg}], "usage": usage})


    run = Run(page_text="x")
    models = ModelSettings()
    http = httpx.Client(transport=httpx.MockTransport(fake_deepseek), base_url=models.deepseek_base_url)
    llm = DeepSeekLLM(models, api_key="pretend-key", cost_sink=run.add_cost, client=http)

    print("== 1. What gets sent ==")
    replies[:] = ['{"topics": ["file sharing", "client approvals"]}']
    out = llm.structured("List the topics of this page.", "Share files with clients.", Topics)
    body = sent[-1]
    print("result:", out)
    print("model:", body["model"], "| response_format:", body["response_format"],
          "| temperature:", body.get("temperature"))
    system = body["messages"][0]["content"]
    print("system prompt ends with:", system[system.index("Reply with"):][:110], "...")

    replies[:] = ['{"topics": ["a"]}']
    llm.structured("s", "u", Topics, tier="judgment")
    print("judgment tier -> model", sent[-1]["model"], "| temperature sent?",
          "temperature" in sent[-1])

    print("\n== 2. Cost of one call ==")
    miss, hit, out_price = models.llm_prices["deepseek-chat"]
    by_hand = (800 * miss + 200 * hit + 500 * out_price) / 1_000_000
    print(f"(800 x {miss} + 200 x {hit} + 500 x {out_price}) / 1,000,000 = ${by_hand:.6f}")
    print("logged:", [(c.label, c.usd) for c in run.costs])

    print("\n== 3. Invalid JSON, then valid ==")
    replies[:] = ['{"topic": "oops"}', '{"topics": ["fixed"]}']
    print("result:", llm.structured("s", "u", Topics))
    retry = sent[-1]["messages"]
    print("the retry adds", len(retry) - 2, "messages; the last one starts:",
          repr(retry[-1]["content"][:40]))

    replies[:] = ["not json", "{}"]
    try:
        llm.structured("s", "u", Topics)
    except LLMOutputError as exc:
        print("twice invalid ->", type(exc).__name__ + ":", str(exc).splitlines()[0])

    print("\n== 4. Embeddings ==")
    calls: list[int] = []


    def fake_gemini(request: httpx.Request) -> httpx.Response:
        reqs = json.loads(request.content)["requests"]
        calls.append(len(reqs))
        raw = {"file sharing": [3.0, 4.0], "sharing files": [3.2, 3.9], "banana bread": [4.0, -3.0]}
        return httpx.Response(200, json={"embeddings": [
            {"values": raw[r["content"]["parts"][0]["text"]]} for r in reqs]})


    cache = DailyCache(SCRATCH / "ex07_embed_cache", today=lambda: date(2026, 9, 25))
    shutil.rmtree(cache.root, ignore_errors=True)
    emb = GeminiEmbeddings(models, "pretend-key", cache, cost_sink=run.add_cost,
                           client=httpx.Client(transport=httpx.MockTransport(fake_gemini),
                                               base_url=GEMINI_URL))
    a, b, c, a2 = emb.embed(["file sharing", "sharing files", "banana bread", "file sharing"])
    print("vector for 'file sharing':", a, "length", round(math.hypot(*a), 6))
    print(f"cosine(file sharing, sharing files) = {cosine(a, b):.3f}")
    print(f"cosine(file sharing, banana bread)  = {cosine(a, c):.3f}")
    emb.embed(["banana bread"])
    print("texts sent per server call:", calls, "(the duplicate and the repeat were not sent)")


def ex08_text_helpers() -> None:
    """Chapter 8: the pure-code helpers (text.py, page_types.py, difficulty.py). Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex08_text_helpers

    What it shows:
    1. words(): how the engine splits text into lowercase words.
    2. ngram_candidates(): the most frequent 2 to 4 word phrases in our example page.
    3. passages(): the page cut into chunks of at most N words, paragraphs kept together.
    4. mmr(): picking phrases that are relevant to the page but not copies of each other.
    5. guess_page_type(): the URL and title rules, in order.
    6. computed_difficulty(): free-mode difficulty from the top 10, including the empty-results
       trap (REVIEW.md finding M1).

    Things to try:
    - Change `diversity` in section 4 from 0.5 to 0.0, then to 0.9.
    - Add your own URLs to the list in section 5.
    - Give the "mixed" results in section 6 two more big sites and watch the score climb.
    """

    # `fakes` lives in tests/, which the offline world put on sys.path at the top of the file.
    from fakes import FakeEmbed

    from seo_engine.config import Thresholds
    from seo_engine.difficulty import computed_difficulty, intent_from_types, site_strength
    from seo_engine.page_types import guess_page_type
    from seo_engine.providers.search import SerpItem, SerpResults
    from seo_engine.providers.tranco import DictRanks
    from seo_engine.text import mmr, ngram_candidates, passages, truncate_words, words

    print("== 1. words() ==")
    print(words("Emitii's client-portal: 2 plans, C++ & A/B tests. Only $9.99!"))

    print("\n== 2. ngram_candidates() on the example page (top 12) ==")
    cands = ngram_candidates(OUR_PAGE)
    print(f"{len(cands)} candidates; first 12:", cands[:12])

    print("\n== 3. passages() ==")
    chunks = passages(OUR_PAGE, max_words=40)
    for i, chunk in enumerate(chunks, 1):
        print(f"  [{i}] {len(chunk.split()):>2} words: {chunk[:70]}...")
    print("truncate_words(page, 12):", repr(truncate_words(OUR_PAGE, 12)))

    print("\n== 4. mmr(): relevant AND different ==")
    embed = FakeEmbed()
    doc = "client portal file sharing client approvals"
    options = ["client portal", "client portals", "client portal app", "file sharing",
               "client approvals", "banana bread"]
    doc_vec, *vecs = embed.embed([doc, *options])
    for diversity in (0.0, 0.5, 0.9):
        picked = [options[i] for i in mmr(doc_vec, vecs, top_n=3, diversity=diversity)]
        print(f"  diversity {diversity}: {picked}")

    print("\n== 5. guess_page_type() ==")
    urls = [
        ("https://example.com/", ""),
        ("https://example.com/pricing", ""),
        ("https://example.com/blog/best-client-portals", ""),
        ("https://example.com/blog/asana-vs-trello", ""),
        ("https://example.com/blog/how-to-onboard-clients", ""),
        ("https://example.com/tools/invoice-generator", ""),
        ("https://www.reddit.com/r/agency/comments/1", ""),
        ("https://youtu.be/abc", ""),
        ("https://example.com/whitepaper.pdf", ""),
        ("https://example.com/login", ""),
        ("https://example.com/blog/our-story", "Top 10 tools we love"),
        ("https://example.com/blog/our-story", ""),
    ]
    for url, title in urls:
        extra = f"  (title: {title!r})" if title else ""
        print(f"  {guess_page_type(url, title):<11} {url}{extra}")

    print("\n== 6. computed_difficulty() ==")
    t = Thresholds()
    ranks = DictRanks(TRANCO)


    def serp(*items: tuple[str, str]) -> SerpResults:
        return SerpResults(phrase="x", country="US", items=[
            SerpItem(rank=i, url=u, domain=u.split("/")[2].removeprefix("www."), page_type=k)
            for i, (u, k) in enumerate(items, 1)])


    mixed = serp(
        ("https://www.g2.com/categories/client-portal", "category"),     # rank 900
        ("https://clientflow.com/", "product"),                         # rank 150,000
        ("https://www.reddit.com/r/agency/1", "forum"),                 # forum: 0.1 whatever rank
        ("https://portalpro.io/", "product"),                           # not listed
    )
    for item in mixed.items:
        print(f"  {item.domain:<16} rank {ranks.rank(item.domain)!s:<7} -> strength "
              f"{site_strength(item, ranks, t)}")
    d = computed_difficulty(mixed, ranks, t)
    print(f"  mean {sum(d.strengths)}/{len(d.strengths)} x 100 -> difficulty {d.score}, "
          f"small sites {d.small_sites}")
    print("  intent from these page types:", intent_from_types([i.page_type for i in mixed.items]))

    empty = serp()
    print("  EMPTY results -> difficulty", computed_difficulty(empty, ranks, t).score,
          "(looks easiest of all, see REVIEW.md M1)")

    blog = SerpItem(rank=1, url="https://tiny-agency.medium.com/our-portal",
                    domain="tiny-agency.medium.com")
    platform_ranks = DictRanks({"medium.com": 250})
    print("  a one-person blog on medium.com gets strength",
          site_strength(blog, platform_ranks, t), "(the platform's, see REVIEW.md M4)")


def ex09_keyword_research() -> None:
    """Chapter 9: phrase discovery on its own, three times, on the offline world. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex09_keyword_research

    It calls the real `keyword_research` (src/seo_engine/tools/keyword_research.py) three times:

    1. With default settings (site strength "new"), exactly as chapter 4's run did.
    2. With a stricter difficulty ceiling for a "new" site: 15 instead of 30.
    3. With the empty-SERP trap: we give a broad phrase NO Google results. Its computed
       difficulty becomes 0, so it looks easy and wins (docs/REVIEW.md, finding M1).

    Things to try:
    - Run 1 with site_strength="established" (ceiling 60). Does anything new pass? Why not?
    - In run 3, give the trap phrase a Bing volume of 5 instead of 900. Does it still win?
    - Run 1 with thresholds=Thresholds(seeds_to_expand=1). How many candidates now?
    """

    from seo_engine.config import Thresholds
    from seo_engine.tools.keyword_research import KeywordResearch, keyword_research


    def show(title: str, **settings) -> KeywordResearch:
        run = make_run(**settings)
        deps = make_deps(run)
        out = keyword_research(deps, run.page_text, run.settings, TODAY)
        ceiling = run.settings.thresholds.difficulty_ceiling[run.settings.site_strength]
        print(f"\n=== {title} (site strength {run.settings.site_strength!r}, ceiling {ceiling})")
        print(f"  {len(out.candidates)} candidates, {len(out.clusters)} clusters, "
              f"{len(out.questions)} questions kept as gap evidence")
        print("  Candidates that reached the difficulty step:")
        for c in out.candidates:
            if c.difficulty is not None:
                print(f"    {c.keyword:<42} vol {c.volume:<5} diff {c.difficulty:<4} fit {c.fit:.2f}"
                      f"  {c.reason or '(no reason: joined a cluster as a member)'}")
        print("  Clusters (best first):")
        for cl in out.clusters:
            print(f"    {cl.head:<42} {cl.fit:.4f} x {cl.demand:.4f} x {cl.winnability:.4f}"
                  f" = {cl.score:.4f}  members={cl.members}")
        print(f"  CHOSEN: {[p.text for p in out.phrases]}")
        return out


    first = show("1. Default settings")
    p = first.phrases[0]
    print(f"\n  The main Phrase object the brief will carry:\n    {p!r}")

    strict = Thresholds(difficulty_ceiling={"new": 15, "growing": 45, "established": 60})
    show("2. A stricter ceiling", thresholds=strict)

    # 3. The empty-SERP trap. Add a broad phrase with real demand but no Google results.

    SEEDS.append("client project software")
    KEYWORDS["client project software"] = (900, None)
    SERPS.pop("client project software", None)  # no results at all for this phrase
    show("3. Empty-SERP trap: 'client project software' has no search results")


def ex10_intent() -> None:
    """Chapter 10: read search intent from the kinds of pages Google shows. Free, pure code.

        .venv/bin/python docs/guide_exercises.py ex10_intent

    `intent_verdict` (src/seo_engine/tools/serp_top.py) looks at the page types in a top 10 and
    says whether one type dominates, and whether OUR page type fits. This script feeds it
    several made-up top 10s, then shows the keyword-research version (`intent_from_types`).

    Things to try:
    - Add a mix of 6 listicles and 4 unknowns. Why is the share 100%, not 60%?
    - Make a mix of exactly 55% one type. Which verdict do you get, and why is it not in the
      architecture document?
    - Pass our_type=None. What happens to the flag?
    """

    from seo_engine.difficulty import intent_from_types
    from seo_engine.tools.serp_top import intent_verdict

    MIXES = {
        "7 listicles, 3 products (ours: product)": (["listicle"] * 7 + ["product"] * 3, "product"),
        "4 listicles, 3 products, 3 guides": (["listicle"] * 4 + ["product"] * 3 + ["guide"] * 3,
                                             "product"),
        "5 guides, 5 listicles (no products)": (["guide"] * 5 + ["listicle"] * 5, "product"),
        "8 unknown, 2 products": (["unknown"] * 8 + ["product"] * 2, "product"),
        "11 guides, 9 products (55%)": (["guide"] * 11 + ["product"] * 9, "product"),
        "all 10 products": (["product"] * 10, "product"),
    }

    for name, (types, ours) in MIXES.items():
        v = intent_verdict(types, ours)
        print(f"\n{name}")
        print(f"  mix={v.mix} dominant={v.dominant} share={v.dominant_share} verdict={v.verdict}")
        print(f"  flag: {v.flag}")

    print("\nKeyword research reads a phrase's intent from the same page types:")
    for types in (["guide", "guide", "product"], ["product", "listicle", "tool", "tool"],
                  ["forum", "video", "guide"], ["unknown", "pdf"]):
        print(f"  {types} -> {intent_from_types(types)!r}")


def ex11_competitors() -> None:
    """Chapter 11: the five competitor filters, on the offline world. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex11_competitors

    It gets the Google results for the 3 phrases chapter 4 chose, then calls the real
    `competitor_analysis` (src/seo_engine/tools/competitor_analysis.py) twice:
    1. With default thresholds: 7 pages kept, 6 dropped, each with its reason.
    2. With max_pages_per_domain=1 and competitors_min=8, to watch filter 5 bite harder and
       the "only N competitor pages" note appear.

    Things to try:
    - Set length_ratio_max=20.0. Which page comes back?
    - Remove "g2.com" from Thresholds.authority_domains. What does g2.com get dropped for now?
    - Give make_run(thresholds=Thresholds(competitors_max=4)). Read the new drop reasons.
    """

    from seo_engine.config import Thresholds
    from seo_engine.tools.competitor_analysis import competitor_analysis, pool_results

    PHRASES = ["client project workspace", "client portal for agencies", "file sharing with clients"]


    def show(title: str, **settings) -> None:
        run = make_run(**settings)
        deps = make_deps(run)
        serps = [deps.search.top(p, run.settings.country, run.settings.pages_per_phrase)
                 for p in PHRASES]
        print(f"\n=== {title}")
        pool = pool_results(serps)
        print(f"  {sum(len(s.items) for s in serps)} results for 3 phrases -> "
              f"{len(pool)} unique URLs after pooling (best rank kept)")
        out = competitor_analysis(deps, serps, run.page_text, run.settings)
        print(f"  Our page: type={out.ours.page_type!r}, topics={out.ours.topics}")
        for p in out.kept:
            print(f"  KEPT     {p.source:<10} {p.url}")
        for d in out.dropped:
            print(f"  DROPPED  google#{d.rank:<3} {d.url:<46} {d.reason}")
        print(f"  type mix of pages read: {out.type_mix}")
        print(f"  notes: {out.notes}")


    show("1. Default thresholds")
    show("2. One page per domain, want at least 8",
         thresholds=Thresholds(max_pages_per_domain=1, competitors_min=8))


def ex12_topic_coverage() -> None:
    """Chapter 12: topic coverage, gaps and the score, on the offline world. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex12_topic_coverage

    Steps:
    1. Rebuild chapter 4's competitors (keyword research -> Google results -> competitor filters).
    2. Call the real `topic_coverage` (src/seo_engine/tools/topic_coverage.py).
    3. Recompute the score by hand, term by term, and check it matches `content_score`.
    4. Show how synonyms merge before counting, and how the 90th percentile is worked out.

    Things to try:
    - Add "file sharing" as a line to OUR_PAGE in the offline world (top of this file). Watch the score jump.
    - Pass Thresholds(must_cover_share=0.5) through make_run(thresholds=...). What moves bucket?
    - Change the merge threshold in part 4 to 0.95. Do the two labels still merge?
    """

    from seo_engine.models import Page
    from seo_engine.pipeline import gather_evidence
    from seo_engine.tools.competitor_analysis import competitor_analysis
    from seo_engine.tools.keyword_research import keyword_research
    from seo_engine.tools.topic_coverage import (
        content_score,
        merge_topics,
        percentile,
        saturation,
        topic_coverage,
    )

    run = make_run()
    deps = make_deps(run)
    s = run.settings
    t = s.thresholds

    # 1. The inputs, built exactly as pipeline.py builds them.
    research = keyword_research(deps, run.page_text, s, TODAY)
    phrases = [p.text for p in research.phrases]
    serps = [deps.search.top(p, s.country, s.pages_per_phrase) for p in phrases]
    comp = competitor_analysis(deps, serps, run.page_text, s)
    evidence = gather_evidence(serps, research.questions, [m for p in research.phrases
                                                            for m in p.cluster])
    print(f"{len(comp.kept)} competitors, {len(evidence)} pieces of demand evidence:")
    for e in evidence:
        print(f"   {e.source:<20} {e.text}")

    # 2. The real tool.
    cov = topic_coverage(deps, comp.kept, comp.ours, phrases, evidence, s)
    print(f"\n{'topic':<26}{'bucket':<8}{'covered':<9}{'ours':<6}passages per competitor")
    for c, d in zip(cov.counts, cov.details, strict=True):
        print(f"{c.topic:<26}{c.bucket:<8}{f'{c.covered_by}/{c.total}':<9}{c.ours_passages:<6}"
              f"{d.competitor_passages}  {d.noise_reason}")
    print("\nGaps:")
    for g in cov.gaps:
        print(f"   {g.topic}  (covered by {g.covered_by}; {g.evidence})")

    # 3. The score by hand.
    print("\nScore by hand (must + worth topics only):")
    num = den = 0.0
    for c in cov.counts:
        if c.bucket not in ("must", "worth"):
            continue
        w = c.covered_by / c.total
        sat = saturation(c.ours_passages, t.bm25_k)
        num += w * sat
        den += w * t.score_target_saturation
        print(f"   {c.topic:<24} w = {c.covered_by}/{c.total} = {w:.3f}   "
              f"s = {c.ours_passages}/({c.ours_passages}+{t.bm25_k}) = {sat:.3f}   "
              f"w*s = {w * sat:.3f}   w*0.77 = {w * t.score_target_saturation:.3f}")
    print(f"   sum(w*s) = {num:.4f}, sum(w*0.77) = {den:.4f}, "
          f"100 x {num:.4f} / {den:.4f} = {100 * num / den:.2f} -> {min(100, round(100 * num / den))}")
    print(f"   content_score says: {content_score(cov.counts, t.bm25_k, t.score_target_saturation)[0]}")
    print(f"   saturation for n = 0..5 passages: "
          f"{[round(saturation(n, t.bm25_k), 2) for n in range(6)]}")

    # 4. Synonyms merge, and the percentile used by the stuffing check.
    pages = [
        Page(url="https://a.com/", source="google#1", page_type="product", text="",
             headings=[], topics=["file sharing"]),
        Page(url="https://b.com/", source="google#2", page_type="product", text="",
             headings=[], topics=["Files sharing", "time tracking"]),
    ]
    for g in merge_topics(pages, deps, t.topic_merge_similarity):
        print(f"\nGroup {g.label!r}: members {g.members}, listed by competitors {sorted(g.listed_by)}")
    values = [1, 1, 0, 0, 0, 1]
    print(f"\n90th percentile of {values} = {percentile(values, 0.9)}"
          f"  (sorted {sorted(values)}, position (6-1) x 0.9 = 4.5)")


def ex13_snippet_check() -> None:
    """Chapter 13: how wide is a title on Google? Pure code, free.

        .venv/bin/python docs/guide_exercises.py ex13_snippet_check

    `pixel_width` adds up each character's width in Arial at 20px (Google's desktop title font).
    `snippet_check` then judges a title and description (src/seo_engine/tools/snippet_check.py).

    Things to try:
    - Put your own page's title into MY_TITLE and run again.
    - Compare "iiiiiiiiii" and "WWWWWWWWWW": same length, very different widths.
    - Remove " | Emitii" from a title that fails. Does it pass now?
    """

    from seo_engine.tools.snippet_check import _char_width, pixel_width, snippet_check

    MY_TITLE = "Client Project Workspace for Agencies | Emitii"
    DESCRIPTION = (
        "A client project workspace where agencies share files, collect client approvals and track "
        "tasks. Invite clients in minutes, no email threads."
    )

    print("Character by character, for the word 'Client':")
    total = 0
    for ch in "Client":
        total += _char_width(ch)
        print(f"  {ch!r}: {_char_width(ch)} units   running total {total}")
    print(f"  {total} units x 20px / 1000 = {total * 20 / 1000} -> rounded up: "
          f"{pixel_width('Client')}px")

    print("\nSame length, different width:")
    for text in ("iiiiiiiiii", "WWWWWWWWWW", "Client portal"):
        print(f"  {text!r:<16} {len(text):>2} chars  {pixel_width(text):>4}px")

    print()
    titles = [
        MY_TITLE,
        "Emitii | Client Project Workspace for Agencies",
        "The Complete Client Project Workspace for Marketing Agencies, Studios and Consultancies",
        "Emitii",
    ]
    for title in titles:
        r = snippet_check(title, DESCRIPTION, phrase="client project workspace")
        print(f"{title!r}")
        print(f"  {r.title_chars} chars, {r.title_px}px   title_ok={r.title_ok}   "
              f"description {r.description_chars} chars ok={r.description_ok}   passed={r.passed}")
        for reason in r.reasons:
            print(f"  - {reason}")


def ex14_brief_writer() -> None:
    """Chapter 14: the Brief Writer and the Draft Writer, with a scripted fake LLM. Free.

        .venv/bin/python docs/guide_exercises.py ex14_brief_writer

    1. Prints the exact prompt `write_brief` sends to the LLM (src/seo_engine/brief.py).
    2. The fake LLM's first answer has only a too-long title, so every title fails the snippet
       check and the writer asks for ONE rewrite. The second answer passes.
    3. Shows the 8-item checklist that code (not the LLM) fills in.
    4. Runs `check_draft_content` on a small draft and explains each number.
    5. Runs `write_draft`: target length from the competitors' median, placeholders collected.

    Things to try:
    - Make the second BriefDraft answer too long as well. Does a third call happen?
    - Change the draft's H1 so it lacks the phrase and watch write_draft retry once.
    - Change competitor word counts in step 5 to [300, 400]. What target length is asked for?
    """

    from fakes import FakeLLM
    from seo_engine.brief import (
        DraftOut,
        check_draft_content,
        draft_text,
        phrase_count,
        write_brief,
        write_draft,
    )
    from seo_engine.config import Settings
    from seo_engine.models import Gap, Phrase, TopicCount

    settings = Settings()
    phrases = [
        Phrase(text="client project workspace", volume=40, volume_source="bing", difficulty=13,
               difficulty_source="computed", intent="commercial",
               cluster=["client project workspace", "client workspace app"]),
        Phrase(text="client portal for agencies", volume=90, volume_source="bing", difficulty=28,
               difficulty_source="computed", intent="commercial"),
    ]
    coverage = [
        TopicCount(topic="file sharing", covered_by=7, total=7, ours_passages=0, bucket="must"),
        TopicCount(topic="client approvals", covered_by=5, total=7, ours_passages=1, bucket="must"),
        TopicCount(topic="pricing plans", covered_by=4, total=7, ours_passages=0, bucket="worth"),
        TopicCount(topic="time tracking", covered_by=1, total=7, ours_passages=0, bucket="rare"),
    ]
    gaps = [Gap(topic="Is a client portal secure?", covered_by=0,
                evidence="People Also Ask: Is a client portal secure?")]

    TOO_LONG = "The Complete Client Project Workspace for Marketing Agencies, Studios and Consultancies"
    GOOD = "Client Project Workspace for Agencies | Emitii"
    DESC = ("A client project workspace where agencies share files, collect client approvals and "
            "track tasks. Invite clients in minutes, no email threads.")
    answers = [
        {"titles": [TOO_LONG], "description": DESC, "headings": ["Client project workspace"]},
        {"titles": [GOOD, "Emitii | Client Project Workspace for Agencies"], "description": DESC,
         "headings": ["The client project workspace for agencies", "Share files with clients",
                      "Pricing plans", "Is a client portal secure?"]},
    ]
    prompts: list[str] = []


    def brief_draft(system: str, user: str) -> dict:
        prompts.append(user)
        return answers.pop(0)


    llm = FakeLLM({"BriefDraft": brief_draft})
    result = write_brief(llm, OUR_PAGE, phrases, coverage, gaps, 11, "(arithmetic)", None, [],
                         settings)

    print("=== 1. What the Brief Writer sees (first 1,100 characters of the user message)")
    print(prompts[0][:1100] + " ...")
    print("\n=== 2. Calls made:", llm.calls, " rewrites:", result.rewrites)
    print("    The retry message ends with:\n   ", prompts[1].split("\n\n")[-1].replace("\n", "\n    "))
    print("\n=== 3. The brief")
    print("  titles:", result.brief.titles)
    for snip in result.snippets:
        print(f"  {snip.title!r}: {snip.title_px}px passed={snip.passed} reasons={snip.reasons}")
    print("  must_cover:", [t.topic for t in result.brief.must_cover], "(worth and rare left out)")
    print("  checklist:")
    for key, ok in result.brief.checklist.items():
        print(f"    {'PASS' if ok else 'FAIL'}  {key}")

    print("\n=== 4. Checking a small draft")
    draft = DraftOut(
        h1="The client project workspace for agencies",
        intro="Emitii is a client project workspace where agencies share files and approvals.",
        sections=[{"heading": "Share files", "body": "Upload once. [ADD: storage limit]"},
                  {"heading": "Approvals", "body": "Clients approve on the work."}],
        faq=[{"question": "Is a client portal secure?", "answer": "[ADD: security details]"}],
        cta="Start a free workspace.",
    )
    text = draft_text(draft)
    words = len(text.split())
    uses = phrase_count(text, "client project workspace")
    limit = settings.thresholds.draft_max_density
    print(f"  {words} words; phrase used {uses} times; density = 100 x {uses} x 3 words / {words}"
          f" = {100 * uses * 3 / words:.1f} per 100 words (limit {limit})")
    for c in check_draft_content(draft, "client project workspace", settings):
        print(f"    {'PASS' if c.ok else 'FAIL'}  {c.label}  {c.detail}")

    print("\n=== 5. write_draft")
    draft_prompts: list[str] = []


    def draft_out(system: str, user: str) -> dict:
        draft_prompts.append(user)
        return draft.model_dump()


    final = write_draft(FakeLLM({"DraftOut": draft_out}), OUR_PAGE, result.brief, coverage,
                        [900, 1200, 1800, 2600, 400], settings)
    target = [ln for ln in draft_prompts[0].splitlines() if ln.startswith("TARGET LENGTH")][0]
    print(f"  competitor words [900, 1200, 1800, 2600, 400] -> {target}")
    print(f"  word_count={final.word_count} placeholders={final.placeholders}")
    print(f"  last check: {final.checks[-1].label}: {final.checks[-1].detail}")


def ex15_api_offline() -> None:
    """Chapter 15: drive the real FastAPI backend offline, the way the web app does. Free, no keys.

        .venv/bin/python docs/guide_exercises.py ex15_api_offline

    It builds the real app with `create_app` (src/seo_engine/api/app.py) but hands it the
    offline world instead of real providers. FastAPI's TestClient then sends HTTP requests to
    the app in memory: no server, no port, no network.

    Things to try afterwards:
    - Send a request with "phrases_per_run": 7 and read the 422 error FastAPI returns.
    - Open cache/guide_scratch/runs/ and look at the JSON file the store wrote.
    - Open cache/guide_scratch/report.docx in LibreOffice or Word.
    """

    import io
    import shutil
    import warnings

    warnings.filterwarnings("ignore", message=".*httpx2.*")  # a harmless notice from Starlette

    from docx import Document
    from fastapi.testclient import TestClient

    from seo_engine.api.app import create_app
    from seo_engine.models import Run

    RUNS = SCRATCH / "runs"
    shutil.rmtree(RUNS, ignore_errors=True)  # start with an empty history each time


    def offline_deps(run: Run):
        """Stands in for deps.from_env: the API calls this once per run."""
        run.settings.cache_dir = SCRATCH / "cache"
        return make_deps(run)


    app = create_app(runs_dir=RUNS, deps_factory=offline_deps, web_dist=SCRATCH / "no-web-build")
    client = TestClient(app)

    print("GET /api/health")
    print("  ", client.get("/api/health").json())

    print("GET /api/settings/defaults")
    print("  ", client.get("/api/settings/defaults").json())

    print("POST /api/runs with too little text")
    bad = client.post("/api/runs", json={"page_text": "too short"})
    print("  ", bad.status_code, bad.json()["detail"][0]["msg"])

    print("POST /api/runs with the example page")
    resp = client.post("/api/runs", json={"page_text": OUR_PAGE, "settings": {"country": "US"}})
    print("  ", resp.status_code, resp.json())
    run_id = resp.json()["id"]

    # TestClient runs background tasks before it returns, so the run is already finished here.
    print(f"GET /api/runs/{run_id}")
    rec = client.get(f"/api/runs/{run_id}").json()
    print("   status:", rec["status"], " error:", rec["error"])
    for step in rec["steps"]:
        print(f"   {step['name']:<12} {step['status']:<5} {step['detail']}")
    print("   competitor text sent to the browser?", any("text" in c for c in rec["run"]["competitors"]))
    print("   score:", rec["run"]["brief"]["score"], " titles:", rec["run"]["brief"]["titles"])

    print("GET /api/runs")
    for summary in client.get("/api/runs").json():
        print("  ", summary)

    print(f"GET /api/runs/{run_id}/report.docx")
    report = client.get(f"/api/runs/{run_id}/report.docx")
    path = SCRATCH / "report.docx"
    path.write_bytes(report.content)
    print("  ", report.status_code, report.headers["content-disposition"], f"{len(report.content)} bytes")
    doc = Document(io.BytesIO(report.content))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    print("   headings:", headings[:9])

    print("GET /api/runs/doesnotexist")
    missing = client.get("/api/runs/doesnotexist")
    print("  ", missing.status_code, missing.json())

    print(f"DELETE /api/runs/{run_id}")
    print("  ", client.delete(f"/api/runs/{run_id}").status_code)
    print("   then GET:", client.get(f"/api/runs/{run_id}").status_code)


def ex15_url_guard() -> None:
    """Chapter 15: the URL guard that stops "Import from a website" reading private addresses.

        .venv/bin/python docs/guide_exercises.py ex15_url_guard

    `is_public_url` (src/seo_engine/providers/base.py; it lived in app.py before chapter 20)
    looks the host name up in DNS and refuses any address that is not globally routable:
    private, loopback (this machine), link-local, shared (100.64/10), reserved or multicast.
    An IPv4 address written as IPv6 (::ffff:127.0.0.1) is judged by its IPv4 part.
    `normalise_url` adds "https://" when the user types a bare domain.

    It fetches nothing. It does do DNS lookups, so a public host shows False when you are offline.

    Things to try afterwards:
    - Add "http://0.0.0.0/" and "http://224.0.0.1/". Which rule blocks each one?
    - Turn off your Wi-Fi and run it again. What happens to emitii.com, and why is that safe?
    """

    from seo_engine.api.app import normalise_url
    from seo_engine.providers.base import is_public_url

    TYPED = [
        "emitii.com",
        "https://example.com/pricing",
        "http://localhost:8420/api/health",
        "http://127.0.0.1/",
        "http://10.0.0.5/admin",
        "http://192.168.1.1/",
        "http://169.254.169.254/latest/meta-data/",
        "http://[::1]/",
        "http://100.100.100.200/",
        "http://[::ffff:127.0.0.1]/",
        "not a url at all",
    ]

    for typed in TYPED:
        url = normalise_url(typed)
        verdict = "public, allowed" if is_public_url(url) else "BLOCKED"
        print(f"  {typed:<44} -> {url:<46} {verdict}")


def ex20_site_reader_and_robots() -> None:
    """Chapter 20: robots.txt the RFC 9309 way, which sitemap URLs are skipped, and how 30 pages are picked.
    Free, no network: Gurzu's real robots.txt and sitemap, recorded on 28 Sep 2026.

        .venv/bin/python docs/guide_exercises.py ex20_site_reader_and_robots

    What it shows:
    1. The old parser (Python's urllib.robotparser) and the new one (Protego) disagree on
       Gurzu's real robots.txt. RFC 9309 says the longest matching rule wins and * and $ work.
    2. skip_reason() on real sitemap URLs: a broken Calendly link, a Google verification file,
       a thank-you page, a tag page.
    3. select_pages(): 30 pages spread over sections, so 40 blog posts cannot crowd out the
       service pages.

    Things to try:
    - Add "Disallow: /services/" to the robots text and see which answers change.
    - Call select_pages with n=10 and count the sections.
    """

    from collections import Counter
    from urllib.robotparser import RobotFileParser

    from protego import Protego

    from seo_engine.config import GapSettings
    from seo_engine.providers.sitemap import parse_sitemap, select_pages, skip_reason

    fixtures = REPO / "tests" / "fixtures" / "sitemap"
    robots = (fixtures / "gurzu_robots.txt").read_text()

    print("== 1. robots.txt: old parser vs RFC 9309 ==")
    old = RobotFileParser()
    old.parse(robots.splitlines())
    new = Protego.parse(robots)
    for url in ["https://gurzu.com/admin/", "https://gurzu.com/blog?page=2",
                "https://gurzu.com/services/rails-maintenance/", "https://gurzu.com/services/"]:
        print(f"  {url:<48} old allows: {old.can_fetch('GurzuSEOEngine', url)!s:<5}  "
              f"RFC 9309 allows: {new.can_fetch(url, 'GurzuSEOEngine')}")
    print("  sitemaps named in robots.txt:", list(new.sitemaps))

    print("\n== 2. Which sitemap URLs are skipped ==")
    settings = GapSettings()
    for url in ["https://gurzu.com/https:/calendly.com/gurzu/meeting-with-gurzu",
                "https://gurzu.com/google385b42146547b16e.html",
                "https://gurzu.com/design-ebook-downloaded-thank-you/",
                "https://gurzu.com/tags/rails/", "https://blog.gurzu.com/post",
                "https://gurzu.com/services/web-development/"]:
        print(f"  {url:<66} -> {skip_reason(url, 'gurzu.com', settings) or 'kept'}")

    print("\n== 3. Picking 30 pages ==")
    # The same two filters as SiteReader._usable: skip rules, then robots.txt.
    entries = [e for e in parse_sitemap((fixtures / "gurzu_sitemap.xml").read_bytes()).entries
               if skip_reason(e.url, "gurzu.com", settings) is None
               and new.can_fetch(e.url, "GurzuSEOEngine")]
    in_sitemap = Counter(e.url.split("/")[3] for e in entries)
    picked = select_pages(entries, "https://gurzu.com/", 30)
    chosen = Counter(u.split("/")[3] or "home" for u in picked)
    print("  usable URLs by section:", dict(in_sitemap.most_common()))
    print("  picked by section:     ", dict(chosen.most_common()))


def ex20_numbers_by_hand() -> None:
    """Chapter 20: the Keyword Gap numbers worked by hand: click rates, the Google estimate, visits, lift, categories.
    Free, no network.

        .venv/bin/python docs/guide_exercises.py ex20_numbers_by_hand

    What it shows:
    1. The click-rate table, and the rule that a lower position never gets more clicks.
    2. Google searches = Bing searches x (Google share / Bing share), for the US.
    3. Visits at a position, and the traffic lift of moving from #12 to #4.
    4. Semrush's six categories for a few position patterns.
    5. "Competitor proof": the sum of the competitors' click rates, used to order keywords.

    Things to try:
    - Change the country to "GB" and see the ratio change.
    - Find a position pattern that is in no category at all.
    """

    from seo_engine.config import GapSettings
    from seo_engine.tools.keyword_gap import categories

    s = GapSettings()

    print("== 1. Click rate by position ==")
    raw = s.ctr_by_position
    for pos in (1, 2, 3, 7, 10, 11, 14, 20, 21):
        raw_value = raw[pos - 1] if pos <= len(raw) else 0
        print(f"  #{pos:<3} published {raw_value:.4f}  used {s.ctr_at(pos):.4f}")

    print("\n== 2. From Bing to a rough Google number ==")
    google, bing = s.search_shares["US"]
    ratio = s.google_per_bing("US")
    print(f"  US: Google {google}% / Bing {bing}% = {ratio:.3f}")
    print(f"  100 Bing searches -> about {round(100 * ratio)} Google searches")

    print("\n== 3. Visits and traffic lift ==")
    estimate = round(100 * ratio)
    print(f"  at #4:  {estimate} x {s.ctr_at(4)} = {round(estimate * s.ctr_at(4))} visits a month")
    print(f"  at #12: {estimate} x {s.ctr_at(12)} = {round(estimate * s.ctr_at(12))} visits a month")
    lift = round(estimate * (s.ctr_at(4) - s.ctr_at(12)))
    print(f"  lift from #12 to #4: {estimate} x ({s.ctr_at(4)} - {s.ctr_at(12)}) = {lift}")

    print("\n== 4. Categories (you, then two competitors) ==")
    for ours, theirs in [(None, [3, 8]), (None, [3, None]), (15, [3, 8]), (2, [3, 8]),
                         (5, [3, 8]), (5, [None, None]), (None, [None, None])]:
        print(f"  you {str(ours):<5} competitors {str(theirs):<12} -> {categories(ours, theirs)}")

    print("\n== 5. Competitor proof ==")
    for positions in ([3, 8], [15], [1], [11, 12, 13]):
        proof = sum(s.ctr_at(p) for p in positions)
        print(f"  competitors at {str(positions):<13} proof {proof:.4f}")


def ex20_keyword_gap_offline() -> None:
    """Chapter 20: a whole Keyword Gap analysis, offline, step by step, with the CSV it produces.
    Free, no network: the fake sites, LLM and Google from tests/gap_fakes.py.

        .venv/bin/python docs/guide_exercises.py ex20_keyword_gap_offline

    What it shows:
    1. The five steps and the progress each one reports (the web app shows these).
    2. The keywords found, their demand and business fit, and the brand keyword left out.
    3. Positions and categories for every keyword, then the top keywords to add with reasons.
    4. The first lines of the CSV download.

    Things to try:
    - In tests/gap_fakes.py, move ours.com above moxo.com for "client portal
      software" in SERPS and run again. Which category does it land in now?
    """

    from gap_fakes import deps, gap_run

    from seo_engine.gap_pipeline import run_gap
    from seo_engine.gap_report import keywords_csv

    run = gap_run()

    def on_step(name: str, status: str, detail: str) -> None:
        if status == "done" or name == "google":
            print(f"  [{name:<8}] {status:<7} {detail}")

    print("== 1. The steps ==")
    run_gap(run, deps(), on_step)

    print("\n== 2. Keywords found ==")
    for k in run.discovery.keywords:
        print(f"  {k.keyword:<30} bing {str(k.bing_searches):<5} fit {k.business_fit} sites {sorted(k.pages)}")
    print("  left out as brand keywords:", [k.keyword for k in run.discovery.brand_keywords])

    print("\n== 3. Positions and categories ==")
    for row in run.result.rows:
        print(f"  {row.keyword:<30} {row.positions}  {row.categories}")
    print("  top keywords to add:")
    for o in run.result.top:
        print(f"    {o.keyword}: {o.reason}")

    print("\n== 4. The CSV ==")
    for line in keywords_csv(run).splitlines()[:3]:
        print("  " + line[:110])


def ex21_site_snapshot_offline() -> None:
    """Chapter 21: a whole Site Snapshot, offline: the steps, the facts, the position groups, the checks and the CSV.
    Free, no network: the fake site, LLM, Google and fact sources from tests/snapshot_fakes.py.

        .venv/bin/python docs/guide_exercises.py ex21_site_snapshot_offline

    What it shows:
    1. The five steps. The facts started in the background at the beginning, so the facts step
       only collects them.
    2. Every fact with its source, and what "no data" looks like.
    3. The searches, their positions, the position groups and the visits sum.
    4. The ten technical checks.
    5. A snapshot where the fact sources fail: notes, not a failed snapshot.

    Things to try:
    - In tests/snapshot_fakes.py, move ours.com to #1 for "secure client portal" in SERPS.
      How do the groups and the visits change?
    - Run it with `site="blog.ours.com"`. Why are the Majestic and Tranco facts empty?
    """
    import sqlite3

    from snapshot_fakes import FakeProber, snapshot_deps

    from seo_engine.snapshot_pipeline import SnapshotRun, run_snapshot
    from seo_engine.snapshot_report import snapshot_csv

    def on_step(name: str, status: str, detail: str) -> None:
        if status == "done":
            print(f"  [{name:<8}] done    {detail}")

    print("== 1. The steps ==")
    run = SnapshotRun(site="ours.com")
    run_snapshot(run, snapshot_deps(), on_step)
    r = run.result

    print("\n== 2. Facts ==")
    f = r.facts
    print(f"  link score {f.link.score} / 10, {f.link.referring_domains} referring domains, {len(f.link.history)} months")
    print(f"  Majestic {f.majestic.ref_subnets} linking networks; Tranco rank {f.tranco_rank}")
    print(f"  speed: {[(v.metric, v.p75, v.status) for v in r.vitals]}")
    print(f"  registered {f.dates.registered}, first seen online {f.dates.first_seen}")
    print(f"  sitemap: {r.sitemap_urls} pages in {r.sitemap_files} file(s)")

    print("\n== 3. Searches and positions ==")
    for k in r.keywords:
        print(f"  {k.keyword:<28} position {str(k.position):<5} visits {k.visits}")
    print("  groups:", [(g.label, g.count) for g in r.groups])
    print(f"  visits a month: {r.visits} (from {r.visits_keywords} searches with Bing numbers)")
    print("  competitors:", [(c.domain, c.keywords) for c in r.competitors])

    print("\n== 4. Technical checks ==")
    for c in r.checks.checks:
        print(f"  {c.status:<7} {c.label}")
    print("  CSV:", snapshot_csv(run).splitlines()[1])

    print("\n== 5. When the fact sources fail ==")

    class BrokenList:
        def lookup(self, domain: str, exact: bool = False):
            raise sqlite3.OperationalError("database or disk is full")

    bad = SnapshotRun(site="ours.com")
    run_snapshot(bad, snapshot_deps(majestic=BrokenList(), prober=FakeProber(fail=True)))
    print("  status: finished, with notes:", [n for n in bad.notes if "failed" in n])


def ex21_site_checks_by_hand() -> None:
    """Chapter 21: the technical check rules by hand: noindex headers, redirect kinds, robots.txt answers.
    Free, no network: tools/site_checks.py on made-up answers.

        .venv/bin/python docs/guide_exercises.py ex21_site_checks_by_hand

    What it shows:
    1. Which X-Robots-Tag headers mean "noindex" for Google, and which do not.
    2. Permanent (301, 308) and temporary (302, 307) redirects.
    3. How Google reads each robots.txt answer, and what we do ourselves (the page fetcher's
       stricter rules).

    Things to try:
    - Add the header "googlebot: noindex, otherbot: index" to part 1. Is it noindex for Google?
    - Why do 401 and 403 give "pass" in part 3 but "we do not read"?
    """
    from protego import Protego

    from seo_engine.providers.fetcher import robots_body
    from seo_engine.providers.site_probe import Hop, Probe, SiteProbe
    from seo_engine.tools.site_checks import check_permanent_redirects, check_robots, says_noindex

    print("== 1. X-Robots-Tag ==")
    for headers in (
        ["noindex"],
        ["googlebot: noindex"],
        ["otherbot: noindex"],
        ["max-image-preview:large, noindex"],
        ["otherbot: nofollow", "noindex"],
    ):
        print(f"  {str(headers):<40} noindex for Google: {says_noindex(headers, header=True)}")

    print("\n== 2. Redirects ==")
    home = "https://www.site.com/"
    for code in (301, 308, 302, 307):
        moved = Probe(url="http://site.com/", status=200, final_url=home, hops=[Hop(url="http://site.com/", status=code)])
        probe = SiteProbe(domain="site.com", variants=[moved], home=home)
        c = check_permanent_redirects(probe)
        print(f"  HTTP {code}: {c.status:<5} {c.detail[:70]}")

    print("\n== 3. robots.txt answers ==")
    for status, body in ((200, "User-agent: *\nDisallow: /\n"), (404, ""), (403, ""), (503, "")):
        probe = SiteProbe(domain="site.com", variants=[], home=home, robots=Probe(url="r", status=status, body=body))
        google = check_robots(probe).status
        rules, unreachable = robots_body(status, body, False)
        we_read = Protego.parse(rules).can_fetch(home, "GurzuSEOEngine")
        print(f"  HTTP {status}: Google check {google:<5} | we read the site: {we_read} (temporary: {unreachable})")


# ========================================================================================
# 4. Chapter 17: your first test file (pytest collects the test_* functions)
# ========================================================================================
# Chapter 17: your first tests. Run them with pytest, not with python:
#
#     .venv/bin/python docs/guide_exercises.py ex17_my_first_test
#
# Each function whose name starts with `test_` is one test. pytest calls it; the test passes
# if no `assert` fails and nothing raises. These tests use the real engine code and the same
# fakes as tests/ (FakeLLM), so they are free and take well under a second.
#
# Things to try afterwards:
# - Break a test on purpose: change 400 to 401 in test_pixel_width_of_a_known_title, run
#   again, and read how pytest shows the difference.
# - Run only one test: add  -k rewrite  to the command.
# - Add a test of your own for `competition()` in src/seo_engine/report.py (30 is "Low",
#   31 is "Medium").

import pytest
from fakes import FakeLLM

from seo_engine.brief import write_brief
from seo_engine.config import Settings
from seo_engine.models import Phrase
from seo_engine.tools.snippet_check import pixel_width, snippet_check

GOOD_DESC = (
    "A client project workspace where agencies share files, collect client approvals and "
    "track tasks. Invite clients in minutes, no email threads."
)


def test_pixel_width_of_a_known_title() -> None:
    # The same title the offline run produced (chapter 4): 400 px at Arial 20 px.
    assert pixel_width("Client project workspace for Agencies | Emitii") == 400


def test_long_title_fails_and_says_why() -> None:
    title = "The Complete Client Project Workspace for Marketing Agencies, Studios and Freelancers"
    result = snippet_check(title, GOOD_DESC, phrase="client project workspace")
    assert not result.title_ok
    assert any("px" in reason for reason in result.reasons)


@pytest.mark.parametrize("n_chars, ok", [(69, False), (70, True), (158, True), (159, False)])
def test_description_length_edges(n_chars: int, ok: bool) -> None:
    # Thresholds.description_min_chars = 70, description_max_chars = 158
    result = snippet_check("Client Project Workspace for Agencies | Emitii", "x" * n_chars)
    assert result.description_ok is ok


def test_brief_writer_rewrites_once_when_every_title_fails() -> None:
    too_long = "The Complete Client Project Workspace for Marketing Agencies and Studios Everywhere"
    good = "Client Project Workspace for Agencies | Emitii"
    replies = [
        {"titles": [too_long], "description": GOOD_DESC, "headings": ["Client project workspace"]},
        {"titles": [good], "description": GOOD_DESC, "headings": ["Client project workspace"]},
    ]
    llm = FakeLLM({"BriefDraft": lambda system, user: replies.pop(0)})
    phrase = Phrase(text="client project workspace", volume=40, difficulty=13, intent="commercial")

    result = write_brief(llm, "page text", [phrase], [], [], 0, "", None, [], Settings())

    assert result.rewrites == 1
    assert llm.calls == ["BriefDraft", "BriefDraft"]  # asked twice: first draft, then the fix
    assert result.brief.titles == [good]
    assert result.brief.checklist["title_width_ok"]


# ========================================================================================
# 5. The command line, and the guide checker
# ========================================================================================
GUIDE = REPO / "docs" / "SEO-Advisor-Guide.md"
EXERCISES = {name: fn for name, fn in sorted(globals().items()) if re.fullmatch(r"ex\d\d_\w+", name)}
TEST_EXERCISE = "ex17_my_first_test"
BLOCK = re.compile(
    r"<!-- exercise:(\w+) -->\n<details><summary>.*?</summary>\n\n```python\n(.*?)```\n\n</details>\n<!-- /exercise -->",
    re.S,
)


def chapter_of(name: str) -> int:
    return int(name[2:4])


def summary(name: str) -> str:
    if name == TEST_EXERCISE:
        return "Your first pytest file: snippet checks and a scripted brief writer"
    return (inspect.getdoc(EXERCISES[name]) or "").splitlines()[0]


def all_names() -> list[str]:
    return sorted([*EXERCISES, TEST_EXERCISE])


def run_one(name: str) -> int:
    print(f"\n##### {name}: {summary(name)}\n")
    if name == TEST_EXERCISE:
        import pytest

        return pytest.main([__file__, "-q", "-p", "no:cacheprovider"])
    EXERCISES[name]()
    return 0


def exercise_source(name: str) -> str:
    """The code a reader sees in the guide for one exercise (or a helper section)."""
    src = Path(__file__).read_text(encoding="utf-8")
    if name == "offline_world":
        return src.split("# 1. The offline world", 1)[1].split("\n", 2)[2].split("# " + "=" * 88 + "\n# 2.")[0].rstrip() + "\n"
    if name == "narrate":
        return src.split("# 2. narrate()", 1)[1].split("\n", 2)[2].split("# " + "=" * 88 + "\n# 3.")[0].rstrip() + "\n"
    if name == TEST_EXERCISE:
        return src.split("# 4. Chapter 17", 1)[1].split("\n", 2)[2].split("# " + "=" * 88 + "\n# 5.")[0].rstrip() + "\n"
    return inspect.getsource(EXERCISES[name])


def sync_guide() -> int:
    text = GUIDE.read_text(encoding="utf-8")

    def repl(m: re.Match) -> str:
        name = m.group(1)
        head = m.group(0).split("\n```python\n", 1)[0]
        return f"{head}\n```python\n{exercise_source(name)}```\n\n</details>\n<!-- /exercise -->"

    new = BLOCK.sub(repl, text)
    GUIDE.write_text(new, encoding="utf-8")
    print("guide updated" if new != text else "guide already in sync")
    return 0


def slug(heading: str) -> str:
    s = re.sub(r"[^\w\- ]", "", heading.strip().lower())
    return s.replace(" ", "-")


def check() -> int:
    problems: list[str] = []
    print("Exercises (each in its own process):")
    for name in all_names():
        if "live" in name:
            print(f"  skip  {name} (live, costs money)")
            continue
        result = subprocess.run(
            [sys.executable, __file__, name], cwd=REPO, capture_output=True, text=True, timeout=300
        )
        print(f"  {'ok' if result.returncode == 0 else 'FAIL':<5} {name}")
        if result.returncode:
            problems.append(f"{name} exited {result.returncode}: {result.stderr[-400:]}")

    text = GUIDE.read_text(encoding="utf-8")
    print(f"Guide: {GUIDE.relative_to(REPO)} ({len(text.splitlines())} lines)")
    ref = re.compile(r"`((?:src|web|tests|scripts|docs|evals)/[\w./-]+?\.\w+):(\d+)(?:-(\d+))?")
    for m in ref.finditer(text):
        path, first, last = REPO / m.group(1), int(m.group(2)), int(m.group(3) or m.group(2))
        if not path.is_file():
            problems.append(f"{m.group(0)}` names a missing file")
            continue
        n = len(path.read_text(encoding="utf-8", errors="replace").splitlines())
        if not 1 <= first <= last <= n:
            problems.append(f"{m.group(0)}` is outside the file ({n} lines)")
    prose = re.sub(r"```.*?```|`[^`\n]*`", "", text, flags=re.S)
    anchors = {slug(h) for h in re.findall(r"^#+ (.+)$", re.sub(r"```.*?```", "", text, flags=re.S), re.M)}
    for target in re.findall(r"\]\(([^)\s]+)\)", prose):
        if "://" in target:
            continue
        if target.startswith("#"):
            if target[1:] not in anchors:
                problems.append(f"link to missing heading {target}")
        elif not (GUIDE.parent / target.split("#")[0]).exists():
            problems.append(f"link to missing file {target}")
    if text.count("—") + text.count("–"):
        problems.append("em or en dashes in the guide")
    no_code = re.sub(r"```.*?```", "", text, flags=re.S)
    chapters = re.split(r"^# (?=Chapter \d+:)", no_code, flags=re.M)[1:]
    for ch in chapters:
        title = ch.splitlines()[0]
        if not title.startswith("Chapter 19:") and not all(s in ch for s in ("## Recap", "## Check yourself")):
            problems.append(f"{title}: missing Recap or Check yourself")
    for m in BLOCK.finditer(text):
        if m.group(2) != exercise_source(m.group(1)):
            problems.append(f"exercise code for {m.group(1)} is out of date: run sync-guide")
    shown = {m.group(1) for m in BLOCK.finditer(text)}
    for name in all_names():
        if name not in shown:
            problems.append(f"{name} is not shown in the guide")

    if problems:
        print(f"\n{len(problems)} problem(s):")
        for p in problems:
            print(f"  - {p}")
        return 1
    print("\nAll good.")
    return 0


def main(argv: list[str]) -> int:
    arg = argv[0] if argv else "list"
    if arg == "list":
        for name in all_names():
            print(f"  {name:<30} {summary(name)}")
        print("\nRun one:  .venv/bin/python docs/guide_exercises.py <name or chapter number>")
        return 0
    if arg == "check":
        return check()
    if arg == "sync-guide":
        return sync_guide()
    if arg == "all":
        return check_all()
    if arg.isdigit():
        chosen = [n for n in all_names() if chapter_of(n) == int(arg)]
        if not chosen:
            print(f"Chapter {arg} has no exercises. Try: list")
            return 1
        code = 0
        for name in chosen:
            if "live" in name:
                print(f"\n(skipping {name}: it costs money; run it by name)")
                continue
            code |= run_one(name)
        return code
    if arg in all_names():
        return run_one(arg)
    print(f"Unknown exercise {arg!r}. Try: list")
    return 1


def check_all() -> int:
    code = 0
    for name in all_names():
        if "live" in name:
            continue
        result = subprocess.run([sys.executable, __file__, name], cwd=REPO)
        code |= result.returncode
    return code


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
